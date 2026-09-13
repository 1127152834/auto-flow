#!/usr/bin/env python3
"""Measure the PM2 Excel path with a real temporary database and files."""

from __future__ import annotations

import hashlib
import json
import sys
from datetime import UTC, datetime, timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
from time import perf_counter, sleep
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "apps" / "backend"))

from tests.contract.test_settings_dashboard import _app

ROWS = 10_000
PAGE_SIZE = 200
REPORT = ROOT / "docs/project-management/implementation/pm2-data-volume-verification.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def elapsed(call):
    started = perf_counter()
    value = call()
    return value, round(perf_counter() - started, 6)


def register(client, project: str, path: Path, purpose: str, proof: str) -> str:
    token = str(uuid4())
    response = client.post(
        "/internal/project-files/selections",
        headers={
            "x-autoflow-host-token": "host",
            "x-autoflow-file-window-token": proof,
        },
        json={
            "selectionToken": token,
            "path": str(path),
            "projectId": project,
            "windowId": 7,
            "purpose": purpose,
            "expiresAt": (datetime.now(UTC) + timedelta(minutes=5)).isoformat(),
        },
    )
    assert response.status_code == 204, response.text
    return token


def operation(client, project: str, key: str) -> dict:
    for _ in range(2_000):
        response = client.get(
            f"/api/v1/projects/{project}/operations/by-idempotency-key/{key}"
        )
        assert response.status_code == 200, response.text
        value = response.json()
        if value["status"] in {"succeeded", "failed"}:
            return value
        sleep(0.01)
    raise TimeoutError(f"operation {key} did not settle")


def main() -> None:
    with TemporaryDirectory(prefix="autoflow-pm2-volume-") as directory:
        temp = Path(directory)
        source, target = temp / "source.xlsx", temp / "export.xlsx"
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("Records")
        sheet.append(["编号", "名称", "分组"])
        for index in range(1, ROWS + 1):
            sheet.append([f"{index:06d}", f"记录 {index}", f"组 {index % 17:02d}"])
        workbook.save(source)
        source_before = sha256(source)

        app = _app(temp)
        with TestClient(
            app, headers={"x-autoflow-token": "renderer"}
        ) as client:
            project = client.post(
                "/api/v1/projects",
                headers={"Idempotency-Key": str(uuid4())},
                json={"name": "PM2 volume verification"},
            ).json()["projectId"]
            proof = str(uuid4())
            public_headers = {
                "x-autoflow-file-window-id": "7",
                "x-autoflow-file-window-token": proof,
            }
            input_token = register(client, project, source, "inspectExcel", proof)
            inspect_key = str(uuid4())

            inspect_response, inspect_seconds = elapsed(
                lambda: client.post(
                    f"/api/v1/projects/{project}/table-imports/excel/inspect",
                    headers={**public_headers, "Idempotency-Key": inspect_key},
                    json={"selectionToken": input_token},
                )
            )
            assert inspect_response.status_code == 200, inspect_response.text
            inspection = inspect_response.json()["inspection"]
            assert inspection["sheets"][0]["rowCount"] == ROWS

            definitions = [
                ("code", "编号"),
                ("name", "名称"),
                ("group", "分组"),
            ]
            import_key = str(uuid4())
            def run_import():
                response = client.post(
                    f"/api/v1/projects/{project}/table-imports/excel",
                    headers={**public_headers, "Idempotency-Key": import_key},
                    json={
                        "name": "十千行验证",
                        "inspectionId": inspection["inspectionId"],
                        "fingerprint": inspection["fingerprint"],
                        "sheetId": inspection["sheets"][0]["sheetId"],
                        "mapping": [
                            {
                                "columnIndex": index,
                                "target": {
                                    "kind": "new",
                                    "definition": {
                                        "key": key,
                                        "name": name,
                                        "type": "string",
                                        "required": True,
                                        "validation": {},
                                    },
                                },
                            }
                            for index, (key, name) in enumerate(definitions)
                        ],
                        "identity": {"mode": "column", "columnIndex": 0},
                    },
                )
                assert response.status_code == 202, response.text
                return operation(client, project, import_key)

            imported, import_seconds = elapsed(run_import)
            assert imported["status"] == "succeeded", imported
            table = imported["result"]["table"]
            assert table["recordCount"] == ROWS
            base = f"/api/v1/projects/{project}/tables/{table['tableId']}"

            def read_pages():
                count, first, last, pages = 0, None, None, 0
                while count < ROWS:
                    response = client.get(
                        base + "/records",
                        params={
                            "datasetGeneration": table["datasetGeneration"],
                            "page": pages + 1,
                            "pageSize": PAGE_SIZE,
                        },
                    )
                    assert response.status_code == 200, response.text
                    items = response.json()["items"]
                    assert items
                    first = first or items[0]["ref"]["recordKey"]["value"]
                    last = items[-1]["ref"]["recordKey"]["value"]
                    count += len(items)
                    pages += 1
                return count, pages, first, last

            page_result, pagination_seconds = elapsed(read_pages)
            assert page_result == (ROWS, 50, "000001", "010000"), page_result
            fields = client.get(base + "/fields").json()["items"]
            output_token = register(client, project, target, "exportXlsx", proof)
            export_key = str(uuid4())
            def run_export():
                response = client.post(
                    base + "/exports/xlsx",
                    headers={**public_headers, "Idempotency-Key": export_key},
                    json={
                        "selectionToken": output_token,
                        "datasetGeneration": table["datasetGeneration"],
                        "scope": "all",
                        "fieldIds": [item["ref"]["fieldId"] for item in fields],
                        "includeStatus": False,
                    },
                )
                assert response.status_code == 202, response.text
                return operation(client, project, export_key)

            exported, export_seconds = elapsed(run_export)
            assert exported["status"] == "succeeded", exported

        output = load_workbook(target, read_only=True, data_only=True).active
        values = output.iter_rows(values_only=True)
        header = next(values)
        first = next(values)
        last, exported_rows = first, 1
        for row in values:
            last, exported_rows = row, exported_rows + 1
        source_after = sha256(source)
        assert exported_rows == ROWS
        assert first[0] == "000001" and last[0] == "010000"
        assert source_before == source_after
        export_hash = sha256(target)
        assert exported["result"]["sha256"] == export_hash

        report = {
            "schemaVersion": 1,
            "measuredAt": datetime.now(UTC).isoformat(),
            "command": "uv run python ../../scripts/measure-pm2-data.py (from apps/backend)",
            "scope": {
                "rows": ROWS,
                "columns": 3,
                "transport": "FastAPI TestClient",
                "database": "temporary SQLite",
                "uiPerformanceExecuted": False,
            },
            "timingsSeconds": {
                "inspect": inspect_seconds,
                "importPostAndBackgroundWork": import_seconds,
                "getAllRecordPages": pagination_seconds,
                "exportPostAndBackgroundWork": export_seconds,
            },
            "evidence": {
                "inspectionRowCount": inspection["sheets"][0]["rowCount"],
                "importedRecordCount": imported["result"]["importedRecordCount"],
                "pagination": {
                    "pageSize": PAGE_SIZE,
                    "pageCount": page_result[1],
                    "recordCount": page_result[0],
                },
                "sourceSha256Before": source_before,
                "sourceSha256After": source_after,
                "sourceHashPreserved": source_before == source_after,
                "exportSha256": export_hash,
                "exportResultHashMatchesFile": True,
                "exportedDataRowCount": exported_rows,
                "exportedColumnCount": len(header),
                "leadingZeroValues": {"first": first[0], "last": last[0]},
            },
            "limitations": [
                "This is one local backend measurement, not a general large-table performance guarantee.",
                "Renderer/UI responsiveness and memory use were not executed or measured.",
            ],
        }
        REPORT.parent.mkdir(parents=True, exist_ok=True)
        REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
        print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
