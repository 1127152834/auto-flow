from uuid import uuid4

from fastapi.testclient import TestClient
from sqlalchemy import select

from autoflow.infrastructure.database.project_data_models import (
    DataChangeRow,
    DataTableRow,
)
from tests.contract.test_pm2_excel_inspection_recovery import prepare
from tests.contract.test_settings_dashboard import _app


def inspected(client, tmp_path):
    project, token, headers = prepare(client, tmp_path)
    result = client.post(
        f"/api/v1/projects/{project}/table-imports/excel/inspect",
        json={"selectionToken": token},
        headers=headers,
    ).json()["inspection"]
    return project, headers, result


def request_for(inspection):
    return {
        "name": "导入资料",
        "inspectionId": inspection["inspectionId"],
        "fingerprint": inspection["fingerprint"],
        "sheetId": inspection["sheets"][0]["sheetId"],
        "mapping": [
            {
                "columnIndex": 0,
                "target": {
                    "kind": "new",
                    "definition": {
                        "key": "code",
                        "name": "编号",
                        "type": "string",
                        "required": True,
                        "validation": {},
                    },
                },
            }
        ],
        "identity": {"mode": "column", "columnIndex": 0},
    }


def test_new_import_is_durable_and_publishes_table_with_null_status(tmp_path):
    app = _app(tmp_path)
    with TestClient(app, headers={"x-autoflow-token": "renderer"}) as client:
        project, proof, inspection = inspected(client, tmp_path)
        headers = {**proof, "Idempotency-Key": str(uuid4())}
        request = request_for(inspection)
        path = f"/api/v1/projects/{project}/table-imports/excel"
        accepted = client.post(path, json=request, headers=headers)
        assert accepted.status_code == 202, accepted.text
        result = client.get(
            f"/api/v1/projects/{project}/operations/by-idempotency-key/{headers['Idempotency-Key']}"
        ).json()
        assert result["status"] == "succeeded", result
        table = result["result"]["table"]
        assert table["recordCount"] == 1
        source = table["source"]
        assert source["filename"] == "source.xlsx"
        assert source["sheetName"] == inspection["sheets"][0]["name"]
        assert source["importedAt"]
        assert "path" not in source
        public = client.get(
            f"/api/v1/projects/{project}/tables/{table['tableId']}"
        ).json()
        assert public["source"] == source
        records = client.get(
            f"/api/v1/projects/{project}/tables/{table['tableId']}/records",
            params={"datasetGeneration": table["datasetGeneration"]},
        ).json()["items"]
        assert records[0]["ref"]["recordKey"] == {"type": "text", "value": "001"}
        assert records[0]["statusId"] is None
        replay = client.post(path, json=request, headers=headers)
        assert replay.json()["operation"]["operationId"] == result["operationId"]
        assert client.get(f"/api/v1/projects/{project}/tables").json()["total"] == 1
        with app.state.session_factory() as session:
            changes = list(
                session.scalars(
                    select(DataChangeRow).where(
                        DataChangeRow.operation_id == result["operationId"]
                    )
                )
            )
            assert len(changes) == 1
            assert changes[0].after["datasetGeneration"] == table["datasetGeneration"]


def test_invalid_whole_column_identity_never_publishes_empty_shell(tmp_path):
    from openpyxl import load_workbook

    app = _app(tmp_path)
    with TestClient(app, headers={"x-autoflow-token": "renderer"}) as client:
        project, token, proof = prepare(client, tmp_path)
        book = load_workbook(tmp_path / "source.xlsx")
        book.active.append(["001"])
        book.save(tmp_path / "source.xlsx")
        inspection = client.post(
            f"/api/v1/projects/{project}/table-imports/excel/inspect",
            json={"selectionToken": token},
            headers=proof,
        ).json()["inspection"]
        headers = {**proof, "Idempotency-Key": str(uuid4())}
        response = client.post(
            f"/api/v1/projects/{project}/table-imports/excel",
            json=request_for(inspection),
            headers=headers,
        )
        assert response.status_code == 202, response.text
        result = client.get(
            f"/api/v1/projects/{project}/operations/by-idempotency-key/{headers['Idempotency-Key']}"
        ).json()
        assert result["status"] == "failed"
        assert client.get(f"/api/v1/projects/{project}/tables").json()["total"] == 0
        with app.state.session_factory() as session:
            assert not list(
                session.scalars(
                    select(DataTableRow).where(DataTableRow.published.is_(True))
                )
            )


def import_first(client, tmp_path):
    project, proof, inspection = inspected(client, tmp_path)
    headers = {**proof, "Idempotency-Key": str(uuid4())}
    response = client.post(
        f"/api/v1/projects/{project}/table-imports/excel",
        json=request_for(inspection),
        headers=headers,
    )
    assert response.status_code == 202, response.text
    result = client.get(
        f"/api/v1/projects/{project}/operations/by-idempotency-key/{headers['Idempotency-Key']}"
    ).json()
    return project, proof, inspection, result["result"]["table"]


def test_reimport_changes_generation_and_does_not_inherit_status(tmp_path):
    app = _app(tmp_path)
    with TestClient(app, headers={"x-autoflow-token": "renderer"}) as client:
        project, proof, inspection, table = import_first(client, tmp_path)
        base = f"/api/v1/projects/{project}/tables/{table['tableId']}"
        field = client.get(base + "/fields").json()["items"][0]["ref"]["fieldId"]
        status = client.post(
            base + "/statuses",
            headers={"Idempotency-Key": str(uuid4())},
            json={
                "name": "已登记",
                "color": "#875739",
                "order": 0,
                "expectedTableRevision": 1,
            },
        )
        assert status.status_code == 201, status.text
        status_id = status.json()["statusId"]
        changed = client.put(
            base + "/records/MDAx/status",
            headers={"Idempotency-Key": str(uuid4())},
            json={
                "datasetGeneration": table["datasetGeneration"],
                "recordKeyType": "text",
                "expectedStatusRevision": 1,
                "statusId": status_id,
            },
        )
        assert changed.status_code == 200, changed.text
        impact = client.post(base + "/imports/excel/impact").json()
        request = {
            key: value
            for key, value in request_for(inspection).items()
            if key != "name"
        }
        request["mapping"] = [
            {"columnIndex": 0, "target": {"kind": "existing", "fieldId": field}}
        ]
        request.update(
            expectedDatasetGeneration=table["datasetGeneration"],
            expectedTableRevision=impact["expectedRevisions"]["tableRevision"],
            impactRevision=impact["impactRevision"],
        )
        key = str(uuid4())
        response = client.post(
            base + "/imports/excel",
            json=request,
            headers={**proof, "Idempotency-Key": key},
        )
        assert response.status_code == 202, response.text
        result = client.get(
            f"/api/v1/projects/{project}/operations/by-idempotency-key/{key}"
        ).json()
        assert result["status"] == "succeeded", result
        new_table = result["result"]["table"]
        assert new_table["datasetGeneration"] != table["datasetGeneration"]
        current = client.get(
            base + "/records",
            params={"datasetGeneration": new_table["datasetGeneration"]},
        ).json()["items"][0]
        assert current["statusId"] is None and current["linkRevision"] == 1
        old = client.get(
            base + "/records", params={"datasetGeneration": table["datasetGeneration"]}
        )
        assert old.status_code == 410


def test_reimport_detects_record_changes_after_impact_without_losing_old_data(tmp_path):
    app = _app(tmp_path)
    with TestClient(app, headers={"x-autoflow-token": "renderer"}) as client:
        project, proof, inspection, table = import_first(client, tmp_path)
        base = f"/api/v1/projects/{project}/tables/{table['tableId']}"
        field = client.get(base + "/fields").json()["items"][0]["ref"]["fieldId"]
        impact = client.post(base + "/imports/excel/impact").json()
        changed = client.put(
            base + "/records/MDAx/status",
            headers={"Idempotency-Key": str(uuid4())},
            json={
                "datasetGeneration": table["datasetGeneration"],
                "recordKeyType": "text",
                "expectedStatusRevision": 1,
                "statusId": None,
            },
        )
        assert changed.status_code == 200, changed.text
        request = {
            key: value
            for key, value in request_for(inspection).items()
            if key != "name"
        }
        request["mapping"] = [
            {"columnIndex": 0, "target": {"kind": "existing", "fieldId": field}}
        ]
        request.update(
            expectedDatasetGeneration=table["datasetGeneration"],
            expectedTableRevision=table["tableRevision"],
            impactRevision=impact["impactRevision"],
        )
        response = client.post(
            base + "/imports/excel",
            json=request,
            headers={**proof, "Idempotency-Key": str(uuid4())},
        )
        assert response.status_code == 409, response.text
        assert (
            client.get(base).json()["datasetGeneration"] == table["datasetGeneration"]
        )


def test_source_change_after_inspection_preserves_zero_visible_tables(tmp_path):
    app = _app(tmp_path)
    with TestClient(app, headers={"x-autoflow-token": "renderer"}) as client:
        project, proof, inspection = inspected(client, tmp_path)
        (tmp_path / "source.xlsx").write_bytes(b"changed")
        key = str(uuid4())
        response = client.post(
            f"/api/v1/projects/{project}/table-imports/excel",
            json=request_for(inspection),
            headers={**proof, "Idempotency-Key": key},
        )
        assert response.status_code == 202, response.text
        result = client.get(
            f"/api/v1/projects/{project}/operations/by-idempotency-key/{key}"
        ).json()
        assert result["status"] == "failed"
        assert result["error"]["code"] == "EXCEL_SOURCE_CHANGED"
        assert client.get(f"/api/v1/projects/{project}/tables").json()["total"] == 0


def test_chunk_interruption_stays_hidden_and_settles_after_shutdown(
    tmp_path, monkeypatch
):
    from openpyxl import load_workbook

    app = _app(tmp_path)
    with TestClient(app, headers={"x-autoflow-token": "renderer"}) as client:
        project, token, proof = prepare(client, tmp_path)
        book = load_workbook(tmp_path / "source.xlsx")
        for index in range(210):
            book.active.append([f"row-{index}"])
        book.save(tmp_path / "source.xlsx")
        inspection = client.post(
            f"/api/v1/projects/{project}/table-imports/excel/inspect",
            json={"selectionToken": token},
            headers=proof,
        ).json()["inspection"]
        repository = app.state.excel_imports.repository
        original = repository.append
        calls = []

        def interrupt(job, rows):
            original(job, rows)
            calls.append(len(rows))
            assert app.state.settings_runtime.gate.pause(list) == [
                "api_mutation_in_progress"
            ]
            # Force shutdown admission stop independently of the active HTTP request.
            app.state.excel_imports.shutdown()

        monkeypatch.setattr(repository, "append", interrupt)
        key = str(uuid4())
        response = client.post(
            f"/api/v1/projects/{project}/table-imports/excel",
            json=request_for(inspection),
            headers={**proof, "Idempotency-Key": key},
        )
        assert response.status_code == 202, response.text
        result = client.get(
            f"/api/v1/projects/{project}/operations/by-idempotency-key/{key}"
        ).json()
        assert result["status"] == "failed"
        assert calls == [200]
        assert client.get(f"/api/v1/projects/{project}/tables").json()["total"] == 0
        assert app.state.excel_imports.pending_operations() == []


def test_import_accepted_before_process_loss_is_not_replayed_at_startup(
    tmp_path, monkeypatch
):
    app = _app(tmp_path)
    with TestClient(app, headers={"x-autoflow-token": "renderer"}) as client:
        project, proof, inspection = inspected(client, tmp_path)
        monkeypatch.setattr(app.state.excel_imports, "run", lambda _: None)
        key = str(uuid4())
        response = client.post(
            f"/api/v1/projects/{project}/table-imports/excel",
            json=request_for(inspection),
            headers={**proof, "Idempotency-Key": key},
        )
        assert response.status_code == 202
    from autoflow.application.project_data import excel_import

    monkeypatch.setattr(
        excel_import,
        "read_sheet",
        lambda *args: (_ for _ in ()).throw(AssertionError("must not reread old file")),
    )
    with TestClient(_app(tmp_path), headers={"x-autoflow-token": "renderer"}) as client:
        result = client.get(
            f"/api/v1/projects/{project}/operations/by-idempotency-key/{key}"
        ).json()
        assert result["status"] == "failed"
        assert result["error"]["code"] == "EXCEL_IMPORT_INTERRUPTED"
        assert client.get(f"/api/v1/projects/{project}/tables").json()["total"] == 0


def test_final_reimport_publish_rechecks_changes_during_staging(tmp_path, monkeypatch):
    app = _app(tmp_path)
    with TestClient(app, headers={"x-autoflow-token": "renderer"}) as client:
        project, proof, inspection, table = import_first(client, tmp_path)
        base = f"/api/v1/projects/{project}/tables/{table['tableId']}"
        field = client.get(base + "/fields").json()["items"][0]["ref"]["fieldId"]
        impact = client.post(base + "/imports/excel/impact").json()
        request = {k: v for k, v in request_for(inspection).items() if k != "name"}
        request["mapping"] = [
            {"columnIndex": 0, "target": {"kind": "existing", "fieldId": field}}
        ]
        request.update(
            expectedDatasetGeneration=table["datasetGeneration"],
            expectedTableRevision=table["tableRevision"],
            impactRevision=impact["impactRevision"],
        )
        original = app.state.excel_imports.repository.append

        def concurrent_write(job, rows):
            original(job, rows)
            changed = client.put(
                base + "/records/MDAx/status",
                headers={"Idempotency-Key": str(uuid4())},
                json={
                    "datasetGeneration": table["datasetGeneration"],
                    "recordKeyType": "text",
                    "expectedStatusRevision": 1,
                    "statusId": None,
                },
            )
            assert changed.status_code == 200, changed.text

        monkeypatch.setattr(
            app.state.excel_imports.repository, "append", concurrent_write
        )
        key = str(uuid4())
        response = client.post(
            base + "/imports/excel",
            json=request,
            headers={**proof, "Idempotency-Key": key},
        )
        assert response.status_code == 202, response.text
        result = client.get(
            f"/api/v1/projects/{project}/operations/by-idempotency-key/{key}"
        ).json()
        assert (
            result["status"] == "failed"
            and result["error"]["code"] == "REVISION_CONFLICT"
        )
        assert (
            client.get(base).json()["datasetGeneration"] == table["datasetGeneration"]
        )
        assert (
            client.get(
                base + "/records",
                params={"datasetGeneration": table["datasetGeneration"]},
            ).json()["items"][0]["statusRevision"]
            == 2
        )


def test_import_rejects_inspection_grant_from_previous_service_instance(tmp_path):
    app = _app(tmp_path)
    with TestClient(app, headers={"x-autoflow-token": "renderer"}) as client:
        project, proof, inspection = inspected(client, tmp_path)
        app.state.excel_imports.repository.instance_id = "another-service"
        response = client.post(
            f"/api/v1/projects/{project}/table-imports/excel",
            json=request_for(inspection),
            headers={**proof, "Idempotency-Key": str(uuid4())},
        )
        assert response.status_code == 410, response.text
        assert response.json()["error"]["code"] == "FILE_SELECTION_EXPIRED"
        assert client.get(f"/api/v1/projects/{project}/tables").json()["total"] == 0


def test_header_only_workbook_publishes_a_real_empty_dataset_and_allows_new_record(
    tmp_path,
):
    from openpyxl import load_workbook

    app = _app(tmp_path)
    with TestClient(app, headers={"x-autoflow-token": "renderer"}) as client:
        project, token, proof = prepare(client, tmp_path)
        book = load_workbook(tmp_path / "source.xlsx")
        book.active.delete_rows(2)
        book.save(tmp_path / "source.xlsx")
        book.close()
        inspection = client.post(
            f"/api/v1/projects/{project}/table-imports/excel/inspect",
            json={"selectionToken": token},
            headers=proof,
        ).json()["inspection"]
        assert inspection["sheets"][0]["rowCount"] == 0
        body = {**request_for(inspection), "identity": {"mode": "system"}}
        key = str(uuid4())
        assert (
            client.post(
                f"/api/v1/projects/{project}/table-imports/excel",
                json=body,
                headers={**proof, "Idempotency-Key": key},
            ).status_code
            == 202
        )
        result = client.get(
            f"/api/v1/projects/{project}/operations/by-idempotency-key/{key}"
        ).json()
        assert result["status"] == "succeeded", result
        table = result["result"]["table"]
        assert table["recordCount"] == 0
        base = f"/api/v1/projects/{project}/tables/{table['tableId']}"
        field = client.get(f"{base}/fields").json()["items"][0]
        created = client.post(
            f"{base}/records",
            json={
                "datasetGeneration": table["datasetGeneration"],
                "values": [{"fieldId": field["ref"]["fieldId"], "value": "001"}],
            },
            headers={"Idempotency-Key": str(uuid4())},
        )
        assert created.status_code == 201, created.text
        assert created.json()["statusId"] is None
        assert client.get(base).json()["recordCount"] == 1


def test_duplicate_excel_rows_keep_independent_state_and_source_bytes(tmp_path):
    from openpyxl import load_workbook

    from autoflow.domain.project_data.identity import RecordKey, encode_record_key

    app = _app(tmp_path)
    with TestClient(app, headers={"x-autoflow-token": "renderer"}) as client:
        project, token, proof = prepare(client, tmp_path)
        source = tmp_path / "source.xlsx"
        book = load_workbook(source)
        book.active.append(["001"])
        book.save(source)
        book.close()
        original_bytes = source.read_bytes()
        inspected_response = client.post(
            f"/api/v1/projects/{project}/table-imports/excel/inspect",
            json={"selectionToken": token}, headers=proof,
        )
        assert inspected_response.status_code == 200, inspected_response.text
        inspection = inspected_response.json()["inspection"]
        body = {**request_for(inspection), "identity": {"mode": "system"}}
        identity = {**proof, "Idempotency-Key": str(uuid4())}
        accepted = client.post(
            f"/api/v1/projects/{project}/table-imports/excel", json=body, headers=identity,
        )
        assert accepted.status_code == 202, accepted.text
        operation = client.get(
            f"/api/v1/projects/{project}/operations/by-idempotency-key/{identity['Idempotency-Key']}"
        ).json()
        assert operation["status"] == "succeeded", operation
        table = operation["result"]["table"]
        base = f"/api/v1/projects/{project}/tables/{table['tableId']}"
        params = {"datasetGeneration": table["datasetGeneration"]}
        rows = client.get(base + "/records", params=params).json()["items"]
        assert len(rows) == 2 and rows[0]["values"] == rows[1]["values"]
        assert rows[0]["ref"]["recordKey"] != rows[1]["ref"]["recordKey"]
        first, second = rows
        url = base + "/records/" + encode_record_key(RecordKey(**first["ref"]["recordKey"]))
        status = client.post(
            base + "/statuses", headers={"Idempotency-Key": str(uuid4())},
            json={"name": "已登记", "color": "#875739", "order": 0,
                  "expectedTableRevision": table["tableRevision"]},
        )
        assert status.status_code == 201, status.text
        changed = client.put(
            url + "/status", headers={"Idempotency-Key": str(uuid4())},
            json={**params, "recordKeyType": "uuid", "statusId": status.json()["statusId"],
                  "expectedFromStatusId": None, "expectedStatusRevision": first["statusRevision"]},
        )
        assert changed.status_code == 200, changed.text
        edited = client.patch(
            url, headers={"Idempotency-Key": str(uuid4())},
            json={**params, "recordKeyType": "uuid", "expectedContentRevision": first["contentRevision"],
                  "values": [{"fieldId": first["values"][0]["fieldId"], "value": "本地新值"}]},
        )
        assert edited.status_code == 200, edited.text
        rows = client.get(base + "/records", params=params).json()["items"]
        assert next(row for row in rows if row["ref"] == second["ref"]) == second
        first = edited.json()
        assert first["statusId"] == status.json()["statusId"]
        created = client.post(
            base + "/records", headers={"Idempotency-Key": str(uuid4())},
            json={**params, "values": [{"fieldId": v["fieldId"], "value": v["value"]} for v in second["values"]]},
        )
        assert created.status_code == 201, created.text
        assert created.json()["statusId"] is None
        assert created.json()["ref"] not in [first["ref"], second["ref"]]
        impact = client.post(
            f"/api/v1/projects/{project}/mutation-impact",
            json={"action": "deleteRecord", "target": {"type": "record", "recordRef": first["ref"]}},
        )
        assert impact.status_code == 200, impact.text
        deleted = client.request(
            "DELETE", url, headers={"Idempotency-Key": str(uuid4())},
            json={**params, "recordKeyType": "uuid", "expectedContentRevision": first["contentRevision"],
                  "expectedStatusRevision": first["statusRevision"], "expectedLinkRevision": first["linkRevision"],
                  "impactRevision": impact.json()["impactRevision"]},
        )
        assert deleted.status_code == 202, deleted.text
        remaining = client.get(base + "/records", params=params).json()["items"]
        assert len(remaining) == 2
        assert next(row for row in remaining if row["ref"] == second["ref"]) == second
        assert all(row["ref"] != first["ref"] for row in remaining)
        assert source.read_bytes() == original_bytes


def test_import_preserves_safe_business_format_error_for_record_diagnosis(tmp_path):
    from openpyxl import load_workbook

    app = _app(tmp_path)
    with TestClient(app, headers={"x-autoflow-token":"renderer"}) as client:
        project, token, proof = prepare(client, tmp_path)
        book = load_workbook(tmp_path / 'source.xlsx')
        book.active['B1'] = '金额'; book.active['B2'] = 'not-a-number'; book.save(tmp_path / 'source.xlsx')
        inspected_result = client.post(f'/api/v1/projects/{project}/table-imports/excel/inspect', json={'selectionToken':token}, headers=proof).json()['inspection']
        body = request_for(inspected_result)
        body['mapping'].append({'columnIndex':1,'target':{'kind':'new','definition':{'key':'amount','name':'金额','type':'number','required':True,'validation':{}}}})
        accepted = client.post(f'/api/v1/projects/{project}/table-imports/excel', json=body, headers={**proof,'Idempotency-Key':str(uuid4())})
        assert accepted.status_code == 202, accepted.text
        operation = client.get(f"/api/v1/projects/{project}/operations/by-idempotency-key/{accepted.json()['operation']['idempotencyKey']}").json()
        assert operation['status'] == 'succeeded', operation
        table = operation['result']['table']
        record = client.get(f"/api/v1/projects/{project}/tables/{table['tableId']}/records", params={'datasetGeneration':table['datasetGeneration']}).json()['items'][0]
        issue = next(item for item in record['validationIssues'] if item['rule'] == 'type')
        assert issue['fieldId']
        assert next(cell['value'] for cell in record['values'] if cell['fieldId'] == issue['fieldId']) == 'not-a-number'
