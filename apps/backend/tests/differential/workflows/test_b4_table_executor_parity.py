from __future__ import annotations

import asyncio
import copy
import importlib
import io
import json
import os
import subprocess
import sys
from datetime import UTC, datetime
from pathlib import Path
from threading import Event
from typing import Any

import pytest
from openpyxl import load_workbook

from autoflow.domain.workflows.execution import (
    BinaryOutputSnapshot,
    ExecutionContext,
    WorkflowClock,
)
from autoflow.infrastructure.filesystem.workflow_artifacts import WorkflowArtifactStore
from autoflow.infrastructure.filesystem.workflow_table_workbook import (
    OpenpyxlTableWorkbookRenderer,
)
from autoflow.providers.browser.workflow_worker import _ThreadCancellation

REPOSITORY_ROOT = Path(__file__).resolve().parents[5]
FROZEN_BACKEND = REPOSITORY_ROOT / "reference" / "WebRPA" / "backend"
FROZEN_HARNESS = Path(__file__).with_name("frozen_table_harness.py")

CLASS_NAMES = (
    "TableAddRowExecutor",
    "TableAddColumnExecutor",
    "TableSetCellExecutor",
    "TableGetCellExecutor",
    "TableDeleteRowExecutor",
    "TableClearExecutor",
    "TableExportExecutor",
)

APPROVED_SOURCE_TYPES = {
    "table_add_row",
    "table_add_column",
    "table_set_cell",
    "table_get_cell",
    "table_delete_row",
    "table_clear",
    "table_export",
}


def _target_executors() -> dict[str, type[Any]]:
    try:
        module = importlib.import_module(
            "autoflow.application.workflows.executors.table"
        )
        executors = [getattr(module, name) for name in CLASS_NAMES]
    except (ImportError, AttributeError) as error:
        pytest.fail(f"table production executor is missing: {error}")
    return {executor().module_type: executor for executor in executors}


def _source_result(payload: dict[str, Any]) -> dict[str, Any]:
    env = dict(os.environ)
    env["PYTHONPATH"] = str(FROZEN_BACKEND)
    completed = subprocess.run(
        [sys.executable, str(FROZEN_HARNESS)],
        input=json.dumps(payload, ensure_ascii=False),
        check=True,
        capture_output=True,
        text=True,
        env=env,
    )
    return json.loads(completed.stdout.splitlines()[-1])


async def _target_result(payload: dict[str, Any]) -> dict[str, Any]:
    context = ExecutionContext(
        variables=copy.deepcopy(payload.get("variables", {})),
        data_rows=copy.deepcopy(payload.get("data_rows", [])),
        current_row=copy.deepcopy(payload.get("current_row", {})),
    )
    try:
        result = await _target_executors()[payload["type"]]().execute(
            copy.deepcopy(payload["config"]), context
        )
        result_payload = {
            "success": result.success,
            "message": result.message,
            "data": result.data,
            "error": result.error,
            "branch": result.branch,
        }
    except Exception as error:  # noqa: BLE001 -- capture frozen exception parity.
        result_payload = {
            "raised": type(error).__name__,
            "error": str(error),
        }
    return {
        "result": result_payload,
        "variables": context.variables,
        "data_rows": context.data_rows,
        "current_row": context.current_row,
    }


CASES: list[dict[str, Any]] = [
    {"type": "table_add_row", "config": {"rowData": '{"a":1,"b":"中"}'}},
    {
        "type": "table_add_row",
        "variables": {"row": '{"a":2}'},
        "data_rows": [{"a": 1}],
        "config": {"rowData": "{row}"},
    },
    {"type": "table_add_row", "config": {"rowData": "[1,2]"}},
    {"type": "table_add_row", "config": {"rowData": "{"}},
    {"type": "table_add_row", "config": {"rowData": ""}},
    {
        "type": "table_add_column",
        "data_rows": [{"a": 1}, {"a": 2, "new": "keep"}],
        "config": {"columnName": "new", "defaultValue": 9},
    },
    {
        "type": "table_add_column",
        "variables": {"column": "status", "default": False},
        "config": {"columnName": "{column}", "defaultValue": "{default}"},
    },
    {"type": "table_add_column", "config": {"columnName": ""}},
    {
        "type": "table_add_column",
        "variables": {"column": ["bad"]},
        "data_rows": [{"a": 1}],
        "config": {"columnName": "{column}"},
    },
    {
        "type": "table_set_cell",
        "data_rows": [{"a": 1}, {"a": 2}],
        "config": {"rowIndex": 0, "columnName": "a", "cellValue": 3},
    },
    {
        "type": "table_set_cell",
        "data_rows": [{"a": 1}, {"a": 2}],
        "config": {"rowIndex": -1, "columnName": "b", "cellValue": "x"},
    },
    {
        "type": "table_set_cell",
        "variables": {"index": 1, "value": 8},
        "data_rows": [{"a": 1}, {"a": 2}],
        "config": {"rowIndex": "{index}", "columnName": "a", "cellValue": "{value}"},
    },
    {
        "type": "table_set_cell",
        "data_rows": [{"a": 1}],
        "config": {"rowIndex": "bad", "columnName": "a"},
    },
    {"type": "table_set_cell", "config": {"rowIndex": 0, "columnName": "a"}},
    {
        "type": "table_set_cell",
        "data_rows": [{"a": 1}],
        "config": {"rowIndex": 2, "columnName": "a"},
    },
    {
        "type": "table_set_cell",
        "data_rows": [{"a": 1}],
        "config": {"rowIndex": 0, "columnName": ""},
    },
    {
        "type": "table_get_cell",
        "data_rows": [{"a": 1}, {"a": None}],
        "config": {"rowIndex": -1, "columnName": "a", "variableName": "out"},
    },
    {
        "type": "table_get_cell",
        "data_rows": [{"a": 1}],
        "config": {"rowIndex": 0, "columnName": "missing", "variableName": "out"},
    },
    {
        "type": "table_get_cell",
        "data_rows": [{"a": 1}],
        "config": {"rowIndex": 0, "columnName": "a"},
    },
    {
        "type": "table_get_cell",
        "data_rows": [{"a": 1}],
        "config": {"rowIndex": -2, "columnName": "a", "variableName": "out"},
    },
    {
        "type": "table_get_cell",
        "config": {"rowIndex": 0, "columnName": "a", "variableName": "out"},
    },
    {
        "type": "table_delete_row",
        "data_rows": [{"a": 1}, {"a": 2}],
        "config": {"rowIndex": -1},
    },
    {"type": "table_delete_row", "data_rows": [{"a": 1}], "config": {"rowIndex": 3}},
    {
        "type": "table_delete_row",
        "data_rows": [{"a": 1}],
        "config": {"rowIndex": "bad"},
    },
    {"type": "table_delete_row", "config": {"rowIndex": 0}},
    {
        "type": "table_clear",
        "data_rows": [{"a": 1}, {"a": 2}],
        "current_row": {"pending": 3},
        "config": {},
    },
    {"type": "table_clear", "config": {}},
    {"type": "table_export", "config": {"exportFormat": "csv"}},
]


@pytest.mark.parametrize(
    "payload",
    CASES,
    ids=[f"{case['type']}-{index}" for index, case in enumerate(CASES)],
)
def test_table_operations_match_frozen_webrpa(payload: dict[str, Any]) -> None:
    assert asyncio.run(_target_result(payload)) == _source_result(payload)


def test_frozen_source_file_has_all_approved_module_types() -> None:
    assert set(_source_result({"operation": "types"})["types"]) == APPROVED_SOURCE_TYPES


def test_target_file_has_all_approved_module_types() -> None:
    assert set(_target_executors()) == APPROVED_SOURCE_TYPES


@pytest.mark.parametrize("module_type", sorted(APPROVED_SOURCE_TYPES))
def test_table_operations_never_require_a_browser(module_type: str) -> None:
    assert _target_executors()[module_type]().requires_browser is False


class _RecordingArtifacts:
    def __init__(
        self,
        *,
        failure: Exception | None = None,
        existing_binary_content: bytes | None = None,
    ) -> None:
        self.failure = failure
        self.text_attempts: list[dict[str, Any]] = []
        self.text_calls: list[dict[str, Any]] = []
        self.byte_calls: list[dict[str, Any]] = []
        self.binary_output_attempts: list[dict[str, Any]] = []
        self.binary_output_calls: list[dict[str, Any]] = []
        self.existing_binary_content = existing_binary_content

    async def write_text(self, **kwargs: Any) -> str:
        self.text_attempts.append(kwargs)
        if self.failure is not None:
            raise self.failure
        self.text_calls.append(kwargs)
        return f"/managed/outputs/{kwargs['output_path']}"

    async def write_bytes(self, **kwargs: Any) -> str:
        if self.failure is not None:
            raise self.failure
        self.byte_calls.append(kwargs)
        return f"/managed/artifacts/{kwargs['name']}"

    async def write_binary_output(self, **kwargs: Any) -> str:
        self.binary_output_attempts.append(kwargs)
        if self.failure is not None:
            raise self.failure
        self.binary_output_calls.append(kwargs)
        return f"/managed/outputs/{kwargs['output_path']}"

    async def read_binary_output(self, **kwargs: Any) -> BinaryOutputSnapshot:
        return BinaryOutputSnapshot(
            content=self.existing_binary_content,
            identity="existing" if self.existing_binary_content is not None else "missing",
        )


class _RecordingWorkbookRenderer:
    def __init__(self, *, content: bytes = b"xlsx", failure: Exception | None = None):
        self.content = content
        self.failure = failure
        self.calls: list[dict[str, Any]] = []

    async def render(self, **kwargs: Any) -> bytes:
        self.calls.append(kwargs)
        if self.failure is not None:
            raise self.failure
        return self.content


def _fixed_clock() -> WorkflowClock:
    return WorkflowClock(now=lambda: datetime(2026, 9, 16, 12, 34, 56, tzinfo=UTC))


@pytest.mark.parametrize(
    ("config", "expected_path"),
    [
        ({"exportFormat": "csv"}, "data_20260916_123456.csv"),
        (
            {"exportFormat": "csv", "fileNamePattern": "report_{时间戳}.xlsx"},
            "report_20260916_123456.csv",
        ),
        (
            {"exportFormat": "csv", "savePath": "reports", "fileNamePattern": "rows"},
            "reports/rows.csv",
        ),
        (
            {"exportFormat": "csv", "savePath": "reports/custom.xls"},
            "reports/custom.csv",
        ),
    ],
)
def test_csv_export_uses_managed_text_artifact(
    config: dict[str, Any], expected_path: str
) -> None:
    artifacts = _RecordingArtifacts()
    context = ExecutionContext(
        data_rows=[
            {"name": "Ada", "active": True, "meta": {"x": 1}},
            {"name": "李雷", "score": 2, "active": False},
        ],
        artifacts=artifacts,
        clock=_fixed_clock(),
    )

    result = asyncio.run(
        _target_executors()["table_export"]().execute(
            {**config, "variableName": "path"}, context
        )
    )

    assert result.success is True
    assert result.data == {
        "path": f"/managed/outputs/{expected_path}",
        "rows": 2,
        "format": "csv",
        "sheet_name": None,
        "file_size": len(artifacts.text_calls[0]["content"].encode("utf-8")),
    }
    assert context.variables["path"] == result.data["path"]
    assert artifacts.byte_calls == []
    assert artifacts.text_calls == [
        {
            "output_path": expected_path,
            "content": 'name,active,meta,score\nAda,true,"{""x"": 1}",\n李雷,false,,2\n',
            "separator": "\n",
            "encoding": "utf-8",
            "append": False,
            "mime_type": "text/csv",
        }
    ]


@pytest.mark.parametrize(
    ("source_payload", "target_config"),
    [
        (
            {
                "operation": "export_contract",
                "data_rows": [{"a": 1}, {"a": 2}],
                "config": {
                    "exportFormat": "csv",
                    "fileNamePattern": "report_{时间戳}.xlsx",
                    "variableName": "path",
                },
            },
            {
                "exportFormat": "csv",
                "fileNamePattern": "report_{时间戳}.xlsx",
                "variableName": "path",
            },
        ),
        (
            {
                "operation": "export_contract",
                "save_mode": "file",
                "save_name": "explicit.xls",
                "data_rows": [{"a": 1}, {"a": 2}],
                "config": {"exportFormat": "csv", "variableName": "path"},
            },
            {
                "exportFormat": "csv",
                "savePath": "explicit.xls",
                "variableName": "path",
            },
        ),
    ],
)
def test_csv_export_name_and_result_contract_match_frozen_webrpa(
    source_payload: dict[str, Any], target_config: dict[str, Any]
) -> None:
    source = _source_result(source_payload)
    artifacts = _RecordingArtifacts()
    context = ExecutionContext(
        data_rows=copy.deepcopy(source_payload["data_rows"]),
        artifacts=artifacts,
        clock=_fixed_clock(),
    )

    target = asyncio.run(
        _target_executors()["table_export"]().execute(target_config, context)
    )
    target_path = target.data["path"]
    target_name = Path(target_path).name

    assert source["success"] is True
    assert source["collector"] == {
        "rows": source_payload["data_rows"],
        "format": "csv",
    }
    assert source["data"] == {
        "path": target_name,
        "rows": 2,
        "format": "csv",
        "sheet_name": None,
        "file_size": None,
    }
    assert source["message"] == target.message.replace(target_path, target_name)
    assert source["variables"] == {"path": target_name}


def test_frozen_excel_contract_records_sheet_and_styled_workbook_branch() -> None:
    source = _source_result(
        {
            "operation": "export_contract",
            "data_rows": [{"a": 1}, {"a": 2}],
            "config": {
                "exportFormat": "excel",
                "fileNamePattern": "report",
                "sheetName": "结果",
                "variableName": "path",
            },
        }
    )

    assert source["success"] is True
    assert source["data"] == {
        "path": "report.xlsx",
        "rows": 2,
        "format": "excel",
        "sheet_name": "结果",
        "file_size": None,
    }
    assert source["collector"] == {
        "rows": [{"a": 1}, {"a": 2}],
        "format": "excel",
        "sheet_name": "结果",
    }
    assert source["variables"] == {"path": "report.xlsx"}


def test_excel_export_uses_renderer_and_managed_binary_output() -> None:
    artifacts = _RecordingArtifacts()
    renderer = _RecordingWorkbookRenderer(content=b"real-xlsx")
    context = ExecutionContext(
        data_rows=[{"a": 1}],
        artifacts=artifacts,
        table_workbooks=renderer,
        clock=_fixed_clock(),
    )

    result = asyncio.run(
        _target_executors()["table_export"]().execute(
            {
                "fileNamePattern": "report",
                "sheetName": "结果",
                "variableName": "path",
            },
            context,
        )
    )

    assert result.success is True
    assert result.data == {
        "path": "/managed/outputs/report.xlsx",
        "rows": 1,
        "format": "excel",
        "sheet_name": "结果",
        "file_size": len(b"real-xlsx"),
    }
    assert result.message == (
        "已导出 1 行数据到: /managed/outputs/report.xlsx (Sheet: 结果)"
    )
    assert context.variables == {"path": "/managed/outputs/report.xlsx"}
    assert renderer.calls == [
        {
            "rows": [{"a": 1}],
            "sheet_name": "结果",
            "existing_content": None,
            "cancellation": None,
        }
    ]
    assert artifacts.binary_output_calls == [
        {
            "output_path": "report.xlsx",
            "content": b"real-xlsx",
                "mime_type": (
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                ),
                "expected_identity": "missing",
            }
        ]
    assert artifacts.text_calls == []
    assert artifacts.byte_calls == []


def test_excel_export_requires_renderer_without_touching_artifacts() -> None:
    artifacts = _RecordingArtifacts()
    context = ExecutionContext(
        data_rows=[{"a": 1}], artifacts=artifacts, clock=_fixed_clock()
    )

    result = asyncio.run(_target_executors()["table_export"]().execute({}, context))

    assert result.success is False
    assert result.error == "Excel导出服务不可用"
    assert artifacts.binary_output_attempts == []


def test_real_excel_renderer_preserves_frozen_workbook_shape() -> None:
    content = asyncio.run(
        OpenpyxlTableWorkbookRenderer().render(
            rows=[
                {"name": "Ada", "meta": {"active": True}, "missing": None},
                {"name": "李雷", "items": [1, 2]},
            ],
            sheet_name="结果",
            existing_content=None,
            cancellation=None,
        )
    )

    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook["结果"]
    assert sheet.freeze_panes == "A2"
    assert [cell.value for cell in sheet[1]] == ["name", "meta", "missing", "items"]
    assert [cell.value for cell in sheet[2]] == [
        "Ada",
        '{"active": true}',
        None,
        None,
    ]
    assert [cell.value for cell in sheet[3]] == ["李雷", None, None, "[1, 2]"]
    assert sheet["A1"].font.bold is True
    assert sheet["A1"].font.color is not None
    assert sheet["A1"].font.color.rgb == "00FFFFFF"
    assert sheet["A1"].fill.fgColor.rgb == "004472C4"
    assert sheet["A2"].fill.fill_type is None
    assert sheet["A3"].fill.fgColor.rgb == "00F2F2F2"
    assert sheet.column_dimensions["A"].width == 8
    workbook.close()


def test_excel_renderer_honors_cancellation_before_materializing_workbook() -> None:
    stopped = Event()
    stopped.set()

    async def render() -> None:
        with pytest.raises(asyncio.CancelledError):
            await OpenpyxlTableWorkbookRenderer().render(
                rows=[{"value": "x"}],
                sheet_name="数据",
                existing_content=None,
                cancellation=_ThreadCancellation(stopped),
            )

    asyncio.run(render())


def test_excel_renderer_preserves_other_sheets_and_replaces_same_name() -> None:
    renderer = OpenpyxlTableWorkbookRenderer()
    first = asyncio.run(
        renderer.render(
            rows=[{"value": "保留"}],
            sheet_name="其他",
            existing_content=None,
            cancellation=None,
        )
    )
    second = asyncio.run(
        renderer.render(
            rows=[{"value": "旧值"}],
            sheet_name="结果",
            existing_content=first,
            cancellation=None,
        )
    )
    third = asyncio.run(
        renderer.render(
            rows=[{"value": "新值"}],
            sheet_name="结果",
            existing_content=second,
            cancellation=None,
        )
    )

    workbook = load_workbook(io.BytesIO(third))
    assert workbook.sheetnames == ["其他", "结果"]
    assert workbook["其他"]["A2"].value == "保留"
    assert workbook["结果"]["A2"].value == "新值"
    assert workbook["结果"]["A2"].fill.fgColor.rgb == "00F2F2F2"
    workbook.close()


def test_table_executor_does_not_import_the_frozen_file_service_or_excel_stack() -> (
    None
):
    module = importlib.import_module("autoflow.application.workflows.executors.table")
    assert module.__file__ is not None
    source = Path(module.__file__).read_text(encoding="utf-8")

    assert "app.services" not in source
    assert "DataCollector" not in source
    assert "openpyxl" not in source
    assert "xlsxwriter" not in source
    assert "polars" not in source


@pytest.mark.parametrize(
    "save_path",
    ["../escape.csv", "C:\\outside\\escape.csv"],
)
def test_export_rejects_paths_outside_the_managed_artifact_root(save_path: str) -> None:
    artifacts = _RecordingArtifacts()
    context = ExecutionContext(
        data_rows=[{"a": 1}], artifacts=artifacts, clock=_fixed_clock()
    )

    result = asyncio.run(
        _target_executors()["table_export"]().execute(
            {"exportFormat": "csv", "savePath": save_path}, context
        )
    )

    assert result.success is False
    assert result.error == "导出路径无效"
    assert artifacts.text_calls == []


def test_csv_export_allows_explicit_absolute_path(tmp_path: Path) -> None:
    artifacts = _RecordingArtifacts()
    context = ExecutionContext(
        data_rows=[{"a": 1}], artifacts=artifacts, clock=_fixed_clock()
    )
    target = tmp_path / "reports" / "out.csv"

    result = asyncio.run(
        _target_executors()["table_export"]().execute(
            {"exportFormat": "csv", "savePath": str(target)}, context
        )
    )

    assert result.success is True
    assert artifacts.text_calls[0]["output_path"] == str(target)


def test_export_rejects_unknown_format_without_touching_artifacts() -> None:
    artifacts = _RecordingArtifacts()
    context = ExecutionContext(
        data_rows=[{"a": 1}], artifacts=artifacts, clock=_fixed_clock()
    )

    result = asyncio.run(
        _target_executors()["table_export"]().execute(
            {"exportFormat": "typo"}, context
        )
    )

    assert result.success is False
    assert result.error == "不支持的导出格式: typo"
    assert artifacts.text_attempts == []
    assert artifacts.binary_output_attempts == []


class _CredentialReader:
    def get_field(self, name: str, field: str) -> Any | None:
        assert (name, field) == ("prod", "password")
        return "S3CRET"


def test_sensitive_table_cell_is_not_logged_or_exported() -> None:
    artifacts = _RecordingArtifacts()
    context = ExecutionContext(
        data_rows=[{"name": "row"}],
        artifacts=artifacts,
        credentials=_CredentialReader(),
        clock=_fixed_clock(),
    )

    set_result = asyncio.run(
        _target_executors()["table_set_cell"]().execute(
            {
                "rowIndex": 0,
                "columnName": "password",
                "cellValue": "{{cred:prod.password}}",
            },
            context,
        )
    )
    serialized = json.dumps(
        {"message": set_result.message, "data": set_result.data}, ensure_ascii=False
    )
    assert set_result.success is True
    assert context.data_rows[0]["password"] == "S3CRET"
    assert "S3CRET" not in serialized

    get_result = asyncio.run(
        _target_executors()["table_get_cell"]().execute(
            {"rowIndex": 0, "columnName": "password", "variableName": "secret"},
            context,
        )
    )
    assert get_result.success is True
    assert get_result.data is None
    assert context.variables["secret"] == "S3CRET"
    assert "secret" in context.sensitive_variables
    assert "S3CRET" not in (get_result.message or "")

    export_result = asyncio.run(
        _target_executors()["table_export"]().execute(
            {"exportFormat": "csv"}, context
        )
    )
    assert export_result.success is False
    assert export_result.error == "数据表格包含凭据值，不能导出"
    assert artifacts.text_attempts == []


def test_add_row_rejects_oversized_json_before_parsing_or_mutation() -> None:
    context = ExecutionContext(data_rows=[{"existing": True}])

    result = asyncio.run(
        _target_executors()["table_add_row"]().execute(
            {"rowData": json.dumps({"value": "x" * (1024 * 1024)})}, context
        )
    )

    assert result.success is False
    assert result.error == "行数据超过工作流安全限制"
    assert context.data_rows == [{"existing": True}]


def test_csv_export_rejects_more_than_500_unique_columns() -> None:
    artifacts = _RecordingArtifacts()
    context = ExecutionContext(
        data_rows=[{f"column-{index}": index for index in range(501)}],
        artifacts=artifacts,
        clock=_fixed_clock(),
    )

    result = asyncio.run(
        _target_executors()["table_export"]().execute(
            {"exportFormat": "csv"}, context
        )
    )

    assert result.success is False
    assert result.error == "CSV导出内容超过工作流安全限制"
    assert artifacts.text_attempts == []


@pytest.mark.parametrize("module_type", ["table_add_column", "table_set_cell"])
def test_table_mutations_reject_a_501st_column_without_partial_commit(
    module_type: str,
) -> None:
    original = {f"column-{index}": index for index in range(500)}
    context = ExecutionContext(data_rows=[original.copy()])
    config = (
        {"columnName": "overflow", "defaultValue": 1}
        if module_type == "table_add_column"
        else {"rowIndex": 0, "columnName": "overflow", "cellValue": 1}
    )

    result = asyncio.run(
        _target_executors()[module_type]().execute(config, context)
    )

    assert result.success is False
    assert result.error == "数据表格列数超过工作流安全限制"
    assert context.data_rows == [original]


def test_excel_renderer_rejects_excessive_cell_count_before_workbook_creation() -> None:
    row = {f"column-{index}": index for index in range(500)}

    async def render() -> None:
        with pytest.raises(ValueError, match="Excel导出单元格数量超过"):
            await OpenpyxlTableWorkbookRenderer().render(
                rows=[row] * 2001,
                sheet_name="数据",
                existing_content=None,
                cancellation=None,
            )

    asyncio.run(render())


def test_excel_renderer_rejects_existing_workbook_over_column_budget() -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.cell(row=1, column=501, value="overflow")
    content = io.BytesIO()
    workbook.save(content)
    workbook.close()

    async def render() -> None:
        with pytest.raises(ValueError, match="已有Excel文件列数超过"):
            await OpenpyxlTableWorkbookRenderer().render(
                rows=[{"a": 1}],
                sheet_name="数据",
                existing_content=content.getvalue(),
                cancellation=None,
            )

    asyncio.run(render())


def test_export_failure_is_atomic_and_does_not_publish_result_variable() -> None:
    artifacts = _RecordingArtifacts(failure=OSError("disk full"))
    context = ExecutionContext(
        data_rows=[{"a": 1}], artifacts=artifacts, clock=_fixed_clock()
    )

    result = asyncio.run(
        _target_executors()["table_export"]().execute(
            {
                "exportFormat": "csv",
                "savePath": "reports/out.csv",
                "variableName": "path",
            },
            context,
        )
    )

    assert result.success is False
    assert result.error == "写入文件失败: disk full（路径: reports/out.csv）"
    assert "path" not in context.variables
    assert len(artifacts.text_attempts) == 1
    assert artifacts.text_calls == []


def test_excel_export_failure_is_atomic_and_does_not_publish_result_variable() -> None:
    artifacts = _RecordingArtifacts(failure=OSError("disk full"))
    renderer = _RecordingWorkbookRenderer(content=b"xlsx")
    context = ExecutionContext(
        data_rows=[{"a": 1}],
        artifacts=artifacts,
        table_workbooks=renderer,
        clock=_fixed_clock(),
    )

    result = asyncio.run(
        _target_executors()["table_export"]().execute(
            {
                "exportFormat": "excel",
                "savePath": "reports/out.xlsx",
                "variableName": "path",
            },
            context,
        )
    )

    assert result.success is False
    assert result.error == "写入文件失败: disk full（路径: reports/out.xlsx）"
    assert "path" not in context.variables
    assert len(artifacts.binary_output_attempts) == 1
    assert artifacts.binary_output_calls == []


def test_export_requires_the_artifact_port() -> None:
    context = ExecutionContext(data_rows=[{"a": 1}], clock=_fixed_clock())

    result = asyncio.run(
        _target_executors()["table_export"]().execute({"exportFormat": "csv"}, context)
    )

    assert result.success is False
    assert (
        result.error
        == "写入文件失败: 文件输出服务不可用（路径: data_20260916_123456.csv）"
    )


def test_export_rejects_more_than_100000_rows_without_calling_writer() -> None:
    artifacts = _RecordingArtifacts()
    context = ExecutionContext(
        data_rows=[{}] * 100_001, artifacts=artifacts, clock=_fixed_clock()
    )

    result = asyncio.run(
        _target_executors()["table_export"]().execute({"exportFormat": "csv"}, context)
    )

    assert result.success is False
    assert result.error == "数据表格超过工作流安全限制"
    assert artifacts.text_calls == []


def test_add_row_rejects_more_than_100000_rows_without_partial_mutation() -> None:
    rows: list[dict[str, Any]] = [{}] * 100_000
    context = ExecutionContext(data_rows=rows.copy())

    result = asyncio.run(
        _target_executors()["table_add_row"]().execute(
            {"rowData": '{"new":1}'}, context
        )
    )

    assert result.success is False
    assert result.error == "数据表格超过工作流安全限制"
    assert context.data_rows == rows


def test_export_rejects_csv_larger_than_8_mib_without_calling_writer() -> None:
    artifacts = _RecordingArtifacts()
    context = ExecutionContext(
        data_rows=[{"value": "x" * (8 * 1024 * 1024)}],
        artifacts=artifacts,
        clock=_fixed_clock(),
    )

    result = asyncio.run(
        _target_executors()["table_export"]().execute({"exportFormat": "csv"}, context)
    )

    assert result.success is False
    assert result.error == "CSV导出内容超过工作流安全限制"
    assert artifacts.text_calls == []


@pytest.mark.parametrize(
    ("module_type", "data_rows", "config"),
    [
        (
            "table_add_column",
            [{"a": index} for index in range(10_000)],
            {"columnName": "new", "defaultValue": 1},
        ),
        (
            "table_export",
            [{"a": index} for index in range(10_000)],
            {"exportFormat": "csv"},
        ),
    ],
)
def test_large_table_operations_cancel_without_partial_commit(
    module_type: str,
    data_rows: list[dict[str, Any]],
    config: dict[str, Any],
) -> None:
    async def cancel_operation() -> tuple[list[dict[str, Any]], _RecordingArtifacts]:
        stopped = Event()
        artifacts = _RecordingArtifacts()
        context = ExecutionContext(
            data_rows=copy.deepcopy(data_rows),
            artifacts=artifacts,
            cancellation=_ThreadCancellation(stopped),
            clock=_fixed_clock(),
        )
        task = asyncio.create_task(
            _target_executors()[module_type]().execute(config, context)
        )
        await asyncio.sleep(0)
        stopped.set()
        with pytest.raises(asyncio.CancelledError):
            await task
        return context.data_rows, artifacts

    final_rows, artifacts = asyncio.run(cancel_operation())
    assert final_rows == data_rows
    assert artifacts.text_calls == []


class _ArtifactRepository:
    def __init__(self) -> None:
        self.calls: list[dict[str, Any]] = []

    def register_artifact(self, **kwargs: Any) -> Any:
        self.calls.append(kwargs)
        return object()


def test_csv_export_uses_the_real_atomic_artifact_writer(tmp_path: Path) -> None:
    repository = _ArtifactRepository()
    store = WorkflowArtifactStore(tmp_path / "workspace", repository)
    writer = store.writer(
        run_id="run-table",
        node_id="table-export",
        execution_id="execution-table",
        purpose="result",
    )
    context = ExecutionContext(
        data_rows=[{"a": 1}, {"a": 2}],
        artifacts=writer,
        clock=_fixed_clock(),
    )

    result = asyncio.run(
        _target_executors()["table_export"]().execute(
            {
                "exportFormat": "csv",
                "savePath": "reports/out.csv",
                "variableName": "path",
            },
            context,
        )
    )

    target = Path(result.data["path"])
    assert result.success is True
    assert (
        target
        == (tmp_path / "workspace/runs/run-table/outputs/reports/out.csv").resolve()
    )
    assert target.read_bytes() == b"a\n1\n2\n"
    assert not list(target.parent.glob(".*.tmp"))
    assert len(repository.calls) == 1
    assert repository.calls[0]["mime_type"] == "text/csv"
    assert repository.calls[0]["size"] == len(target.read_bytes())
    assert context.variables["path"] == str(target)


def test_table_exports_formula_like_text_as_literal_data(tmp_path: Path) -> None:
    repository = _ArtifactRepository()
    store = WorkflowArtifactStore(tmp_path / "workspace", repository)
    writer = store.writer(
        run_id="run-table",
        node_id="table-export",
        execution_id="execution-table",
        purpose="result",
    )
    rows = [
        {
            "=header": "safe",
            "equals": "=1+1",
            "plus": "+SUM(A1:A2)",
            "minus": "-2+3",
            "at": "@SUM(A1:A2)",
        }
    ]

    csv_context = ExecutionContext(
        data_rows=copy.deepcopy(rows), artifacts=writer, clock=_fixed_clock()
    )
    csv_result = asyncio.run(
        _target_executors()["table_export"]().execute(
            {"exportFormat": "csv", "savePath": "reports/literal.csv"},
            csv_context,
        )
    )
    assert csv_result.success is True
    assert Path(csv_result.data["path"]).read_text() == (
        "'=header,equals,plus,minus,at\n"
        "safe,'=1+1,'+SUM(A1:A2),'-2+3,'@SUM(A1:A2)\n"
    )

    xlsx_context = ExecutionContext(
        data_rows=copy.deepcopy(rows),
        artifacts=writer,
        table_workbooks=OpenpyxlTableWorkbookRenderer(),
        clock=_fixed_clock(),
    )
    xlsx_result = asyncio.run(
        _target_executors()["table_export"]().execute(
            {"exportFormat": "excel", "savePath": "reports/literal.xlsx"},
            xlsx_context,
        )
    )
    assert xlsx_result.success is True
    workbook = load_workbook(Path(xlsx_result.data["path"]), data_only=False)
    try:
        sheet = workbook.active
        assert [sheet.cell(1, column).value for column in range(1, 6)] == [
            "=header",
            "equals",
            "plus",
            "minus",
            "at",
        ]
        assert [sheet.cell(1, column).data_type for column in range(1, 6)] == [
            "s",
            "s",
            "s",
            "s",
            "s",
        ]
        assert [sheet.cell(2, column).value for column in range(1, 6)] == [
            "safe",
            "=1+1",
            "+SUM(A1:A2)",
            "-2+3",
            "@SUM(A1:A2)",
        ]
        assert [sheet.cell(2, column).data_type for column in range(1, 6)] == [
            "s",
            "s",
            "s",
            "s",
            "s",
        ]
    finally:
        workbook.close()
