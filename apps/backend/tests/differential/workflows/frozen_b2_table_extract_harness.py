from __future__ import annotations

import asyncio
import io
import json
import sys
import tempfile
from contextlib import redirect_stdout
from pathlib import Path
from typing import Any

from app.executors.base import ExecutionContext
from app.executors.table_extract import ExtractTableDataExecutor
from openpyxl import load_workbook

TABLES = {
    "#headers": [["姓名", "年龄"], ["Ada", "37"], ["李雷", "41"]],
    "#plain": [["Ada", "37"], ["李雷", "41"]],
    "#offset": [["说明", ""], ["姓名", "年龄"], ["Ada", "37"]],
}


class Locator:
    def __init__(
        self, selector: str, rows: list[list[str]], kind: str = "table"
    ) -> None:
        self.selector = selector
        self.rows = rows
        self.kind = kind

    @property
    def first(self) -> Locator:
        return self

    def locator(self, selector: str) -> Locator:
        if selector == "xpath=ancestor-or-self::table":
            return Locator(self.selector, self.rows, "table")
        if selector == "tr":
            return Locator(self.selector, self.rows, "rows")
        if selector == "th, td":
            return Locator(self.selector, self.rows, "cells")
        raise AssertionError(selector)

    def nth(self, index: int) -> Locator:
        if self.kind == "rows":
            return Locator(self.selector, [self.rows[index]], "row")
        if self.kind == "cells":
            return Locator(self.selector, [[self.rows[0][index]]], "cell")
        raise AssertionError(self.kind)

    async def is_visible(self) -> bool:
        return self.selector != "#outside"

    async def evaluate(self, expression: str) -> str:
        if expression == "el => el.tagName":
            return (
                "DIV"
                if self.selector in {"#inside", "#outside"} and self.kind != "table"
                else "TABLE"
            )
        raise AssertionError(expression)

    async def count(self) -> int:
        if self.kind == "rows":
            return len(self.rows)
        if self.kind in {"row", "cells"}:
            return len(self.rows[0])
        return 1

    async def inner_text(self) -> str:
        return self.rows[0][0]


class Page:
    async def wait_for_selector(self, selector: str, **_options: Any) -> None:
        if selector == "#missing":
            raise TimeoutError("missing")

    def locator(self, selector: str) -> Locator:
        rows = TABLES.get(selector, TABLES["#headers"])
        return Locator(selector, rows, "element")


def normalize(result: Any) -> dict[str, Any]:
    data = result.data
    if isinstance(data, dict) and data.get("excelPath"):
        data = {**data, "excelPath": Path(data["excelPath"]).name}
    return {
        "success": result.success,
        "messagePrefix": result.message.split("，", 1)[0],
        "error": result.error,
        "data": data,
    }


async def run(case: str) -> dict[str, Any]:
    context = ExecutionContext(page=Page())
    if case == "headers":
        result = await ExtractTableDataExecutor().execute(
            {"tableSelector": "#headers"}, context
        )
        return {"result": normalize(result), "variables": context.variables}
    if case == "plain":
        result = await ExtractTableDataExecutor().execute(
            {
                "tableSelector": "#plain",
                "includeHeader": False,
                "headerRow": 1,
                "variableName": "rows",
            },
            context,
        )
        return {"result": normalize(result), "variables": context.variables}
    if case == "missing":
        result = await ExtractTableDataExecutor().execute(
            {"tableSelector": "#missing"}, context
        )
        return {"result": normalize(result), "variables": context.variables}
    if case == "outside":
        result = await ExtractTableDataExecutor().execute(
            {"tableSelector": "#outside"}, context
        )
        return {"result": normalize(result), "variables": context.variables}
    if case in {"export_header", "export_plain"}:
        include_header = case == "export_header"
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "table.xlsx"
            result = await ExtractTableDataExecutor().execute(
                {
                    "tableSelector": "#offset",
                    "includeHeader": include_header,
                    "headerRow": 1,
                    "exportToExcel": True,
                    "excelPath": str(target),
                    "variableName": "rows",
                },
                context,
            )
            workbook = load_workbook(target)
            sheet = workbook["表格数据"]
            values = [list(row) for row in sheet.iter_rows(values_only=True)]
            bold = [cell.font.bold for cell in sheet[2]]
            workbook.close()
        return {
            "result": normalize(result),
            "variables": context.variables,
            "values": values,
            "headerBold": bold,
        }
    raise ValueError(case)


if __name__ == "__main__":
    captured = io.StringIO()
    with redirect_stdout(captured):
        result = asyncio.run(run(sys.argv[1]))
    print(json.dumps(result, ensure_ascii=False, separators=(",", ":")))
