from concurrent.futures import Executor, Future
from datetime import UTC, datetime, timedelta
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from autoflow.application.project_data.catalog import DataCatalogService
from autoflow.application.project_data.excel import ProjectExcelService
from autoflow.application.project_data.excel_export import (
    ProjectExcelExportService,
    _request,
)
from autoflow.application.project_data.records import DataRecordService
from autoflow.application.project_data.tables import DataTableService
from autoflow.application.projects.service import ProjectService
from autoflow.domain.projects.models import ProjectError
from autoflow.infrastructure.database.models import Base, ProjectOperationRow
from autoflow.infrastructure.database.project_data import SqlAlchemyProjectData
from autoflow.infrastructure.database.project_data_catalog import (
    SqlAlchemyProjectDataCatalog,
)
from autoflow.infrastructure.database.project_data_models import DataRecordRow
from autoflow.infrastructure.database.project_data_records import (
    SqlAlchemyProjectDataRecords,
)
from autoflow.infrastructure.database.project_excel_exports import (
    SqlAlchemyProjectExcelExports,
)
from autoflow.infrastructure.database.projects import SqlAlchemyProjects
from autoflow.infrastructure.database.session import (
    create_session_factory,
    migrate_database,
)
from autoflow.infrastructure.filesystem.project_excel import write_workbook


def uid():
    return str(uuid4())


class Deferred(Executor):
    def __init__(self):
        self.jobs = []

    def submit(self, fn, /, *args, **kwargs):
        self.jobs.append((fn, args, kwargs))
        return Future()

    def run(self):
        fn, args, kwargs = self.jobs.pop(0)
        fn(*args, **kwargs)


def test_export_accepts_durably_then_recovers_lost_database_completion(
    tmp_path: Path, monkeypatch
):
    database = tmp_path / "db.sqlite3"
    migrate_database(database)
    factory = create_session_factory(database)
    Base.metadata.create_all(factory.kw["bind"])
    project = (
        ProjectService(SqlAlchemyProjects(factory))
        .create(uid(), {"name": "p"})[0]
        .project_id
    )
    table = DataTableService(SqlAlchemyProjectData(factory)).create(
        project, uid(), {"name": "t"}
    )[0]
    field = DataCatalogService(SqlAlchemyProjectDataCatalog(factory)).create_field(
        project,
        table["tableId"],
        uid(),
        {
            "definition": {
                "key": "name",
                "name": "Name",
                "type": "string",
                "required": True,
                "validation": {},
            },
            "expectedTableRevision": 1,
            "sourceColumnPolicy": "localOnly",
        },
    )[0]["field"]
    DataRecordService(SqlAlchemyProjectDataRecords(factory)).create(
        project,
        table["tableId"],
        uid(),
        {
            "datasetGeneration": table["datasetGeneration"],
            "values": [{"fieldId": field["ref"]["fieldId"], "value": "one"}],
        },
    )
    target = tmp_path / "out.xlsx"
    token = uid()
    ProjectExcelService(factory, workspace_id="ws", instance_id="i").register_selection(
        {
            "selectionToken": token,
            "path": str(target),
            "projectId": project,
            "windowId": 1,
            "purpose": "exportXlsx",
            "expiresAt": datetime.now(UTC) + timedelta(minutes=5),
        },
        "w" * 32,
    )
    executor = Deferred()
    repository = SqlAlchemyProjectExcelExports(factory, "ws", "i")
    service = ProjectExcelExportService(repository, None, executor)
    operation = service.start(
        project,
        table["tableId"],
        uid(),
        {
            "selectionToken": token,
            "datasetGeneration": table["datasetGeneration"],
            "scope": "all",
            "fieldIds": [field["ref"]["fieldId"]],
            "includeStatus": False,
        },
        1,
        "w" * 32,
    )
    assert operation.status == "accepted" and not target.exists()
    complete = repository.complete
    monkeypatch.setattr(
        repository,
        "complete",
        lambda _job: (_ for _ in ()).throw(RuntimeError("lost response")),
    )
    executor.run()
    assert target.exists() and service.pending_operations() == [operation.id]
    monkeypatch.setattr(repository, "complete", complete)
    ProjectExcelExportService(repository, None, Deferred()).startup()
    with factory() as session:
        finished = session.get(ProjectOperationRow, operation.id)
        assert finished is not None and finished.status == "succeeded", finished.error
        revision = finished.status_revision
    assert load_workbook(target, read_only=True).active["A2"].value == "one"
    assert len(list(tmp_path.glob("out*.xlsx"))) == 1
    assert service.pending_operations() == []
    with pytest.raises(ProjectError) as wrong_project:
        service.reconcile(uid(), operation.id, uid(), revision)
    assert wrong_project.value.code == "OPERATION_NOT_FOUND"
    with pytest.raises(ProjectError) as old_revision:
        service.reconcile(project, operation.id, uid(), revision - 1)
    assert old_revision.value.code == "OPERATION_REVISION_CONFLICT"
    with pytest.raises(ProjectError) as terminal:
        service.reconcile(project, operation.id, uid(), revision)
    assert terminal.value.code == "OPERATION_NOT_RECONCILABLE"

    target2 = tmp_path / "snapshot.xlsx"
    token2 = uid()
    ProjectExcelService(factory, workspace_id="ws", instance_id="i").register_selection(
        {
            "selectionToken": token2,
            "path": str(target2),
            "projectId": project,
            "windowId": 1,
            "purpose": "exportXlsx",
            "expiresAt": datetime.now(UTC) + timedelta(minutes=5),
        },
        "w" * 32,
    )
    op2 = service.start(
        project,
        table["tableId"],
        uid(),
        {
            "selectionToken": token2,
            "datasetGeneration": table["datasetGeneration"],
            "scope": "all",
            "fieldIds": [field["ref"]["fieldId"]],
            "includeStatus": False,
        },
        1,
        "w" * 32,
    )
    job = repository.claim(op2.id)
    assert job is not None
    snapshot = repository.snapshot(job)
    with factory.begin() as session:
        row = session.scalars(
            select(DataRecordRow).where(DataRecordRow.table_id == table["tableId"])
        ).first()
        assert row
        row.values_json = {field["ref"]["fieldId"]: "two"}
        row.content_revision += 1
    write_workbook(
        target2,
        snapshot.headers,
        snapshot.rows(),
        before_publish=lambda value: repository.prepare(job, value),
    )
    snapshot.close()
    assert load_workbook(target2, read_only=True).active["A2"].value == "one"
    import autoflow.application.project_data.excel_export as export_module

    monkeypatch.setattr(
        export_module,
        "verify_workbook_publication",
        lambda *_: (_ for _ in ()).throw(OSError("unavailable")),
    )
    command = service.reconcile(project, op2.id, uid(), 2)
    assert command.status == "succeeded" and command.result == {
        "targetOperationId": op2.id,
        "status": "failed",
    }
    with factory() as session:
        failed = session.get(ProjectOperationRow, op2.id)
        assert failed and failed.error["code"] == "EXCEL_EXPORT_VERIFICATION_FAILED"
    factory.dispose()


def test_startup_never_replays_an_accepted_export(tmp_path: Path):
    # The production recovery decision is explicit: jobs without publication evidence fail.
    class Repo:
        def pending(self):
            return [type("J", (), {"operation_id": "op", "claim_token": None})()]

        def publication(self, _):
            return None

        def fail(self, operation_id, error):
            self.failure = (operation_id, error)

        def pending_reconciliations(self):
            return []

    repo = Repo()
    service = ProjectExcelExportService(repo, None, Deferred())  # type: ignore[arg-type]
    service.startup()
    assert repo.failure[1]["code"] == "EXCEL_EXPORT_INTERRUPTED"


def test_export_request_requires_a_frozen_filter_and_canonical_fields():
    payload = {
        "selectionToken": uid(),
        "datasetGeneration": uid(),
        "scope": "filter",
        "fieldIds": [],
        "includeStatus": True,
    }
    with pytest.raises(ProjectError):
        _request(payload)
    payload.update(scope="all", fieldIds=["NOT-A-UUID"])
    with pytest.raises(ProjectError):
        _request(payload)
