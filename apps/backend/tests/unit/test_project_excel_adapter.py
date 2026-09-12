from __future__ import annotations

import os
import re
import zipfile
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet._writer import ALL_TEMP_FILES

from autoflow.domain.projects.models import ProjectError
from autoflow.infrastructure.filesystem import project_excel
from autoflow.infrastructure.filesystem.project_excel import (
    inspect_workbook,
    read_sheet,
    write_workbook,
)


def _save_book(path: Path, sheets: dict[str, list[list[object]]]) -> None:
    workbook = Workbook()
    first = True
    for name, rows in sheets.items():
        sheet = workbook.active if first else workbook.create_sheet()
        first = False
        sheet.title = name
        for row in rows:
            sheet.append(row)
    workbook.save(path)
    workbook.close()


def _error_code(error: pytest.ExceptionInfo[ProjectError]) -> str:
    assert error.value.status == 422
    return error.value.code


def test_inspection_preserves_types_and_describes_each_sheet(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    _save_book(
        source,
        {
            "Accounts": [
                ["identity", "active", "created", "score"],
                ["001", True, date(2026, 9, 13), 1.5],
                [None, None, None, None],
                ["002", False, datetime(2026, 9, 13, 8, 30), 2],  # noqa: DTZ001
            ],
            "Empty": [["name"]],
        },
    )

    inspection = inspect_workbook(source)

    assert inspection.filename == "input.xlsx"
    assert len(inspection.fingerprint) == 64
    assert [sheet.name for sheet in inspection.sheets] == ["Accounts", "Empty"]
    assert [sheet.sheet_id for sheet in inspection.sheets] == ["1", "2"]
    assert inspection.sheets[0].headers == ("identity", "active", "created", "score")
    assert inspection.sheets[0].row_count == 2
    assert inspection.sheets[0].ignored_empty_row_count == 1
    assert inspection.sheets[0].sample == (
        (
            "001",
            True,
            {
                "kind": "date",
                "precision": "date",
                "value": "2026-09-13",
                "offset": None,
            },
            1.5,
        ),
        (
            "002",
            False,
            {
                "kind": "date",
                "precision": "datetime",
                "value": "2026-09-13T08:30:00",
                "offset": None,
            },
            2,
        ),
    )
    assert inspection.sheets[1].row_count == 0
    assert inspection.sheets[1].headers == ("name",)
    assert inspection.sheets[1].issues == ()


def test_inspection_sample_is_limited_but_read_sheet_returns_every_row(
    tmp_path: Path,
) -> None:
    source = tmp_path / "many.xlsx"
    _save_book(source, {"Rows": [["value"], *[[index] for index in range(25)]]})

    inspection = inspect_workbook(source)
    with read_sheet(
        source, inspection.sheets[0].sheet_id, inspection.fingerprint
    ) as stream:
        rows = tuple(row.values for row in stream)

    assert len(inspection.sheets[0].sample) == 20
    assert rows == tuple((index,) for index in range(25))


@pytest.mark.parametrize(
    ("rows", "code"),
    [
        ([["name", "name"], ["a", "b"]], "EXCEL_DUPLICATE_HEADER"),
        ([["name", None], ["a", "b"]], "EXCEL_EMPTY_HEADER"),
    ],
)
def test_inspection_rejects_invalid_flat_headers(
    tmp_path: Path, rows: list[list[object]], code: str
) -> None:
    source = tmp_path / "invalid.xlsx"
    _save_book(source, {"Rows": rows})

    with pytest.raises(ProjectError) as caught:
        inspect_workbook(source)

    assert _error_code(caught) == code


def test_inspection_rejects_merged_cells(tmp_path: Path) -> None:
    source = tmp_path / "merged.xlsx"
    _save_book(source, {"Rows": [["group", None], ["name", "email"]]})
    workbook = load_workbook(source)
    workbook.active.merge_cells("A1:B1")
    workbook.save(source)
    workbook.close()

    with pytest.raises(ProjectError) as caught:
        inspect_workbook(source)

    assert _error_code(caught) == "EXCEL_MERGED_CELLS_UNSUPPORTED"


def test_inspection_rejects_formula_without_cached_value(tmp_path: Path) -> None:
    source = tmp_path / "formula.xlsx"
    _save_book(source, {"Rows": [["amount"], ["=1+1"]]})

    with pytest.raises(ProjectError) as caught:
        inspect_workbook(source)

    assert _error_code(caught) == "EXCEL_FORMULA_CACHE_MISSING"


def test_read_sheet_rejects_source_changed_after_inspection(tmp_path: Path) -> None:
    source = tmp_path / "changed.xlsx"
    _save_book(source, {"Rows": [["name"], ["before"]]})
    inspection = inspect_workbook(source)
    _save_book(source, {"Rows": [["name"], ["after"]]})

    with pytest.raises(ProjectError) as caught:
        read_sheet(source, "1", inspection.fingerprint)

    assert _error_code(caught) == "EXCEL_SOURCE_CHANGED"


def test_inspection_rejects_symlink_and_corrupt_archive(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    _save_book(source, {"Rows": [["name"], ["a"]]})
    link = tmp_path / "link.xlsx"
    link.symlink_to(source)

    with pytest.raises(ProjectError) as symlink_error:
        inspect_workbook(link)
    assert _error_code(symlink_error) == "EXCEL_SYMLINK_FORBIDDEN"

    corrupt = tmp_path / "corrupt.xlsx"
    corrupt.write_bytes(b"not an xlsx archive")
    with pytest.raises(ProjectError) as corrupt_error:
        inspect_workbook(corrupt)
    assert _error_code(corrupt_error) == "EXCEL_INVALID_WORKBOOK"


def test_inspection_enforces_actual_zip_and_sheet_limits(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "limited.xlsx"
    _save_book(source, {"Rows": [["a", "b"], [1, 2], [3, 4]]})

    monkeypatch.setattr(
        "autoflow.infrastructure.filesystem.project_excel.MAX_UNCOMPRESSED_BYTES", 100
    )
    with pytest.raises(ProjectError) as zip_error:
        inspect_workbook(source)
    assert _error_code(zip_error) == "EXCEL_ARCHIVE_TOO_LARGE"

    monkeypatch.setattr(
        "autoflow.infrastructure.filesystem.project_excel.MAX_UNCOMPRESSED_BYTES",
        256 * 1024 * 1024,
    )
    monkeypatch.setattr("autoflow.infrastructure.filesystem.project_excel.MAX_ROWS", 1)
    with pytest.raises(ProjectError) as row_error:
        inspect_workbook(source)
    assert _error_code(row_error) == "EXCEL_ROW_LIMIT_EXCEEDED"


def test_inspection_rejects_unsafe_zip_member(tmp_path: Path) -> None:
    source = tmp_path / "unsafe.xlsx"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("../escape.xml", "x")

    with pytest.raises(ProjectError) as caught:
        inspect_workbook(source)

    assert _error_code(caught) == "EXCEL_UNSAFE_ARCHIVE_PATH"


@pytest.mark.parametrize("formula_like", ["=1+1", "+1+1", "-1+1", "@SUM(A1:A2)"])
def test_write_workbook_creates_new_file_with_text_and_dates(
    tmp_path: Path, formula_like: str
) -> None:
    target = tmp_path / "export.xlsx"
    write_workbook(
        target,
        ("identity", "formula-like", "day", "moment"),
        (
            (
                "001",
                formula_like,
                {
                    "kind": "date",
                    "precision": "date",
                    "value": "2026-09-13",
                    "offset": None,
                },
                {
                    "kind": "date",
                    "precision": "datetime",
                    "value": "2026-09-13T08:30:00",
                    "offset": None,
                },
            ),
        ),
    )

    workbook = load_workbook(target, data_only=False)
    sheet = workbook.active
    assert sheet["A2"].value == "001" and sheet["A2"].data_type == "s"
    assert sheet["B2"].value == formula_like and sheet["B2"].data_type == "s"
    assert sheet["C2"].value == datetime(2026, 9, 13)  # noqa: DTZ001
    assert sheet["C2"].number_format == "yyyy-mm-dd"
    assert sheet["D2"].value == "2026-09-13T08:30:00"
    assert sheet["D2"].data_type == "s"
    workbook.close()


def test_write_workbook_refuses_overwrite_and_cleans_temporary_file(
    tmp_path: Path,
) -> None:
    target = tmp_path / "existing.xlsx"
    target.write_bytes(b"original")

    with pytest.raises(ProjectError) as caught:
        write_workbook(target, ("name",), (("value",),))

    assert _error_code(caught) == "EXCEL_OUTPUT_EXISTS"
    assert target.read_bytes() == b"original"
    assert list(tmp_path.iterdir()) == [target]


def test_inspection_binds_open_descriptor_and_rejects_symlink_swap(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "source.xlsx"
    outside = tmp_path / "outside.xlsx"
    _save_book(source, {"Safe": [["name"], ["safe"]]})
    _save_book(outside, {"Secret": [["name"], ["outside"]]})
    original_open = os.open

    def swapping_open(path: Path, flags: int, *args: object, **kwargs: object) -> int:
        source.rename(tmp_path / "held.xlsx")
        source.symlink_to(outside)
        try:
            return original_open(path, flags, *args, **kwargs)
        finally:
            source.unlink()
            (tmp_path / "held.xlsx").rename(source)

    monkeypatch.setattr(
        "autoflow.infrastructure.filesystem.project_excel.os.open", swapping_open
    )
    with pytest.raises(ProjectError) as caught:
        inspect_workbook(source)
    assert _error_code(caught) == "EXCEL_SOURCE_CHANGED"


def test_write_workbook_forces_formula_like_header_to_text(tmp_path: Path) -> None:
    target = tmp_path / "header.xlsx"
    write_workbook(target, ("=1+1",), ())

    workbook = load_workbook(target, data_only=False)
    assert workbook.active["A1"].value == "=1+1"
    assert workbook.active["A1"].data_type == "s"
    workbook.close()


def test_formula_metadata_scans_beyond_sample_and_streams_rows(tmp_path: Path) -> None:
    source = tmp_path / "formulas.xlsx"
    _save_book(
        source,
        {"Rows": [["id", "value"], *[[str(i), i] for i in range(21)], ["x", "=1+1"]]},
    )
    # Give the formula a cached value by editing the worksheet XML directly.
    rewritten = tmp_path / "cached.xlsx"
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(rewritten, "w") as output:
        for entry in original.infolist():
            data = original.read(entry)
            if entry.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb"<f>1\+1</f><v\s*/>", b"<f>1+1</f><v>2</v>", data)
            output.writestr(entry, data)
    rewritten.replace(source)

    inspection = inspect_workbook(source)
    sheet = inspection.sheets[0]
    assert sheet.formula_row_count == (0, 1)
    assert sheet.identity_eligible == (True, False)
    assert len(sheet.sample) == 20

    stream = read_sheet(source, "1", inspection.fingerprint)
    with stream:
        rows = list(stream)
    assert rows[-1].values == ("x", 2)
    assert rows[-1].formula_columns == (1,)
    assert rows[-1].row_number == 23


def test_read_sheet_is_lazy_and_close_releases_workbooks(tmp_path: Path) -> None:
    source = tmp_path / "lazy.xlsx"
    _save_book(source, {"Rows": [["value"], *[[i] for i in range(25)]]})
    inspection = inspect_workbook(source)

    stream = read_sheet(source, "1", inspection.fingerprint)
    with stream:
        first = next(stream)
        assert first.values == (0,)
    assert stream.closed


def test_export_preserves_datetime_offset_as_iso_text(tmp_path: Path) -> None:
    target = tmp_path / "offset.xlsx"
    write_workbook(
        target,
        ("moment",),
        (
            (
                {
                    "kind": "date",
                    "precision": "datetime",
                    "value": "2026-09-13T08:30:00.123456",
                    "offset": "+08:00",
                },
            ),
        ),
    )

    workbook = load_workbook(target, data_only=False)
    cell = workbook.active["A2"]
    assert cell.value == "2026-09-13T08:30:00.123456+08:00"
    assert cell.data_type == "s"
    workbook.close()


def test_shared_strings_limit_uses_content_types_part_name(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "custom-shared.xlsx"
    _save_book(source, {"Rows": [["name"], ["value"]]})
    rewritten = tmp_path / "rewritten.xlsx"
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(rewritten, "w") as output:
        for entry in original.infolist():
            data = original.read(entry)
            if entry.filename == "[Content_Types].xml":
                marker = b"</Types>"
                override = (
                    b'<Override PartName="/xl/customStrings.xml" '
                    b'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
                )
                data = data.replace(marker, override + marker)
            output.writestr(entry, data)
        output.writestr("xl/customStrings.xml", b"0123456789")
    rewritten.replace(source)
    monkeypatch.setattr(
        "autoflow.infrastructure.filesystem.project_excel.MAX_SHARED_STRINGS_BYTES", 5
    )

    with pytest.raises(ProjectError) as caught:
        inspect_workbook(source)
    assert _error_code(caught) == "EXCEL_SHARED_STRINGS_TOO_LARGE"


def test_inspection_rejects_file_changed_while_being_read(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "racing.xlsx"
    _save_book(source, {"Rows": [["name"], ["value"]]})
    original_stat = os.stat(source, follow_symlinks=False)
    changed_stat = list(original_stat)
    changed_stat[8] = original_stat.st_mtime + 1
    calls = 0

    def changing_stat(path: Path) -> os.stat_result:
        nonlocal calls
        calls += 1
        return original_stat if calls == 1 else os.stat_result(changed_stat)

    monkeypatch.setattr(
        "autoflow.infrastructure.filesystem.project_excel.os.lstat", changing_stat
    )
    with pytest.raises(ProjectError) as caught:
        inspect_workbook(source)

    assert _error_code(caught) == "EXCEL_SOURCE_CHANGED"


@pytest.mark.parametrize(
    ("headers", "rows"),
    [(("x" * 32_768,), ()), (("value",), (("😀" * 16_384,),))],
)
def test_write_rejects_strings_beyond_excel_utf16_limit_without_publishing(
    tmp_path: Path, headers: tuple[str, ...], rows: tuple[tuple[str, ...], ...]
) -> None:
    target = tmp_path / "too-long.xlsx"
    with pytest.raises(ProjectError) as caught:
        write_workbook(target, headers, rows)
    assert _error_code(caught) == "EXCEL_STRING_TOO_LONG"
    assert not target.exists()


def test_content_types_is_never_read_without_a_bound(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "bounded.xlsx"
    _save_book(source, {"Rows": [["name"], ["value"]]})
    original_read = zipfile.ZipFile.read

    def guarded_read(
        self: zipfile.ZipFile, name: str, *args: object, **kwargs: object
    ) -> bytes:
        if name == "[Content_Types].xml":
            raise AssertionError("unbounded Content Types read")
        return original_read(self, name, *args, **kwargs)

    monkeypatch.setattr(zipfile.ZipFile, "read", guarded_read)
    project_excel._validate_archive(source.read_bytes(), float("inf"))


def test_stream_close_closes_real_worksheet_sources_and_buffers(tmp_path: Path) -> None:
    source = tmp_path / "resources.xlsx"
    _save_book(source, {"Rows": [["value"], [1], [2]]})
    inspection = inspect_workbook(source)
    stream = read_sheet(source, "1", inspection.fingerprint)
    next(stream)
    row_frame = stream._rows.gi_frame
    assert row_frame is not None
    source_rows = row_frame.f_locals["source_rows"]
    value_rows = row_frame.f_locals["value_rows"]
    source_file = source_rows.gi_frame.f_locals["src"]
    value_file = value_rows.gi_frame.f_locals["src"]
    formula_buffer = stream._formulas._autoflow_buffer
    cached_buffer = stream._cached._autoflow_buffer

    stream.close()

    assert source_file.closed and value_file.closed
    assert formula_buffer.closed and cached_buffer.closed


def test_export_failure_cleans_openpyxl_writer_temporary_files(tmp_path: Path) -> None:
    target = tmp_path / "bad-rows.xlsx"
    before = set(ALL_TEMP_FILES)
    with pytest.raises(ProjectError):
        write_workbook(target, ("value",), ((1,), (2, 3)))
    assert set(ALL_TEMP_FILES) == before
    assert not target.exists()


def test_malformed_numeric_cell_is_project_error_and_closes_stream(
    tmp_path: Path,
) -> None:
    source = tmp_path / "malformed.xlsx"
    _save_book(source, {"Rows": [["value"], [1]]})
    rewritten = tmp_path / "rewritten.xlsx"
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(rewritten, "w") as output:
        for entry in original.infolist():
            data = original.read(entry)
            if entry.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<v>1</v>", b"<v>bogus</v>")
            output.writestr(entry, data)
    rewritten.replace(source)
    fingerprint = __import__("hashlib").sha256(source.read_bytes()).hexdigest()
    stream = read_sheet(source, "1", fingerprint)

    with pytest.raises(ProjectError) as caught:
        next(stream)
    assert _error_code(caught) == "EXCEL_INVALID_WORKBOOK"
    assert stream.closed


@pytest.mark.parametrize("bad_row", [1, 2])
def test_inspection_projects_malformed_header_and_data_cells_to_project_error(
    tmp_path: Path, bad_row: int
) -> None:
    source = tmp_path / f"malformed-row-{bad_row}.xlsx"
    _save_book(
        source, {"Rows": [[1], [2]]} if bad_row == 1 else {"Rows": [["value"], [2]]}
    )
    rewritten = tmp_path / "rewritten.xlsx"
    coordinate = f"A{bad_row}".encode()
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(rewritten, "w") as output:
        for entry in original.infolist():
            data = original.read(entry)
            if entry.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(
                    rb'(<c r="' + coordinate + rb'"[^>]*><v>)[^<]+(</v>)',
                    rb"\1bogus\2",
                    data,
                )
            output.writestr(entry, data)
    rewritten.replace(source)

    with pytest.raises(ProjectError) as caught:
        inspect_workbook(source)
    assert _error_code(caught) == "EXCEL_INVALID_WORKBOOK"
