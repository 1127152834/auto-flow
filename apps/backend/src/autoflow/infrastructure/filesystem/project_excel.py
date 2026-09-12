"""Bounded XLSX inspection, streaming reads, and new-file export."""

from __future__ import annotations

import errno
import hashlib
import io
import math
import os
import stat
import tempfile
import time
import zipfile
from collections.abc import Iterable, Iterator, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path, PurePosixPath
from typing import Any, Self, TypeAlias
from xml.etree.ElementTree import ParseError

from defusedxml.common import DefusedXmlException
from defusedxml.ElementTree import fromstring, iterparse
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles.numbers import is_datetime
from openpyxl.worksheet._writer import ALL_TEMP_FILES
from openpyxl.xml import DEFUSEDXML

from autoflow.domain.projects.models import ProjectError

MAX_FILE_BYTES = 64 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 256 * 1024 * 1024
MAX_SHARED_STRINGS_BYTES = 32 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 10_000
MAX_ROWS = 200_000
MAX_COLUMNS = 500
MAX_CELLS = 10_000_000
MAX_SECONDS = 120.0
SAMPLE_ROWS = 20
DateScalar: TypeAlias = dict[str, str | None]
ScalarValue: TypeAlias = str | int | float | bool | DateScalar | None


@dataclass(frozen=True)
class ExcelRow:
    values: tuple[ScalarValue, ...]
    formula_columns: tuple[int, ...]
    row_number: int


@dataclass(frozen=True)
class SheetInspection:
    sheet_id: str
    name: str
    headers: tuple[str, ...]
    sample: tuple[tuple[ScalarValue, ...], ...]
    row_count: int
    ignored_empty_row_count: int
    formula_row_count: tuple[int, ...]
    identity_eligible: tuple[bool, ...]
    issues: tuple[str, ...] = ()


@dataclass(frozen=True)
class WorkbookInspection:
    fingerprint: str
    filename: str
    sheets: tuple[SheetInspection, ...]


@dataclass
class _Budget:
    cells: int = 0


def _error(code: str, message: str) -> ProjectError:
    return ProjectError(code, message, 422)


def _check(deadline: float) -> None:
    if time.monotonic() > deadline:
        raise _error(
            "EXCEL_TIME_LIMIT_EXCEEDED", "Excel processing exceeded 120 seconds"
        )


def _identity(value: os.stat_result) -> tuple[int, int, int, int]:
    return value.st_dev, value.st_ino, value.st_size, value.st_mtime_ns


def _read_file(path: Path, deadline: float) -> tuple[bytes, str, os.stat_result]:
    if path.suffix.lower() != ".xlsx":
        raise _error(
            "EXCEL_INVALID_WORKBOOK", "The selected file must be an .xlsx workbook"
        )
    descriptor = None
    try:
        before = os.lstat(path)
        if stat.S_ISLNK(before.st_mode):
            raise _error("EXCEL_SYMLINK_FORBIDDEN", "Symbolic links cannot be used")
        if not stat.S_ISREG(before.st_mode):
            raise _error("EXCEL_INVALID_WORKBOOK", "Excel path is not a regular file")
        descriptor = os.open(
            path,
            os.O_RDONLY | getattr(os, "O_BINARY", 0) | getattr(os, "O_NOFOLLOW", 0),
        )
        opened = os.fstat(descriptor)
        if _identity(before) != _identity(opened):
            raise _error("EXCEL_SOURCE_CHANGED", "Excel file changed before open")
        if opened.st_size > MAX_FILE_BYTES:
            raise _error("EXCEL_FILE_TOO_LARGE", "Excel file exceeds 64 MiB")
        chunks = []
        digest = hashlib.sha256()
        total = 0
        while block := os.read(
            descriptor, min(1024 * 1024, MAX_FILE_BYTES + 1 - total)
        ):
            _check(deadline)
            total += len(block)
            if total > MAX_FILE_BYTES:
                raise _error("EXCEL_FILE_TOO_LARGE", "Excel file exceeds 64 MiB")
            chunks.append(block)
            digest.update(block)
        after_fd = os.fstat(descriptor)
        after_path = os.lstat(path)
        if _identity(opened) != _identity(after_fd) or _identity(opened) != _identity(
            after_path
        ):
            raise _error("EXCEL_SOURCE_CHANGED", "Excel file changed while read")
        return b"".join(chunks), digest.hexdigest(), after_fd
    except ProjectError:
        raise
    except OSError as exc:
        if exc.errno in {errno.ELOOP, errno.ENOENT}:
            raise _error(
                "EXCEL_SOURCE_CHANGED", "Excel file changed while opening"
            ) from exc
        raise _error("EXCEL_FILE_UNAVAILABLE", "Excel file cannot be read") from exc
    finally:
        if descriptor is not None:
            os.close(descriptor)


def _safe_name(name: str) -> bool:
    member = PurePosixPath(name)
    return not member.is_absolute() and ".." not in member.parts and "\\" not in name


def _validate_archive(content: bytes, deadline: float) -> None:
    if not DEFUSEDXML:
        raise _error("EXCEL_PARSER_UNAVAILABLE", "Safe XML parser unavailable")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            names = [x.filename for x in entries]
            if len(entries) > MAX_ARCHIVE_ENTRIES or len(names) != len(set(names)):
                raise _error("EXCEL_ARCHIVE_TOO_LARGE", "Archive has too many entries")
            if any(not _safe_name(x) for x in names):
                raise _error("EXCEL_UNSAFE_ARCHIVE_PATH", "Unsafe archive path")
            if any(x.flag_bits & 1 for x in entries):
                raise _error("EXCEL_INVALID_WORKBOOK", "Encrypted workbook unsupported")
            type_parts = []
            type_size = 0
            with archive.open("[Content_Types].xml") as type_stream:
                while block := type_stream.read(
                    min(64 * 1024, 1024 * 1024 + 1 - type_size)
                ):
                    _check(deadline)
                    type_size += len(block)
                    if type_size > 1024 * 1024:
                        raise _error(
                            "EXCEL_ARCHIVE_TOO_LARGE", "Content index too large"
                        )
                    type_parts.append(block)
            types = b"".join(type_parts)
            root = fromstring(types)
            shared = {
                str(x.get("PartName", "")).lstrip("/")
                for x in root
                if str(x.get("ContentType", "")).endswith("sharedStrings+xml")
            }
            total = 0
            for entry in entries:
                part = 0
                with archive.open(entry) as stream:
                    while block := stream.read(1024 * 1024):
                        _check(deadline)
                        part += len(block)
                        total += len(block)
                        if total > MAX_UNCOMPRESSED_BYTES:
                            raise _error(
                                "EXCEL_ARCHIVE_TOO_LARGE",
                                "Expanded workbook exceeds 256 MiB",
                            )
                        if entry.filename in shared and part > MAX_SHARED_STRINGS_BYTES:
                            raise _error(
                                "EXCEL_SHARED_STRINGS_TOO_LARGE",
                                "Shared strings exceed 32 MiB",
                            )
    except ProjectError:
        raise
    except (
        zipfile.BadZipFile,
        KeyError,
        OSError,
        ValueError,
        DefusedXmlException,
        ParseError,
    ) as exc:
        raise _error(
            "EXCEL_INVALID_WORKBOOK", "Excel workbook is invalid or unsupported"
        ) from exc


def _reject_merges(workbook: Any, sheet: Any, deadline: float) -> None:
    try:
        with workbook._archive.open(sheet._worksheet_path) as xml:
            for _, element in iterparse(xml, events=("end",)):
                _check(deadline)
                if element.tag.endswith("}mergeCell"):
                    raise _error(
                        "EXCEL_MERGED_CELLS_UNSUPPORTED", "Merged cells unsupported"
                    )
                element.clear()
    except ProjectError:
        raise
    except (KeyError, OSError, DefusedXmlException, ParseError) as exc:
        raise _error("EXCEL_INVALID_WORKBOOK", "Worksheet XML invalid") from exc


def _scalar(value: Any, number_format: str = "General") -> ScalarValue:
    if isinstance(value, datetime):
        precision = "date" if is_datetime(number_format) == "date" else "datetime"
        return {
            "kind": "date",
            "precision": precision,
            "value": value.date().isoformat()
            if precision == "date"
            else value.isoformat(),
            "offset": None,
        }
    if isinstance(value, date):
        return {
            "kind": "date",
            "precision": "date",
            "value": value.isoformat(),
            "offset": None,
        }
    if isinstance(value, bool) or value is None or isinstance(value, (str, int)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise _error("EXCEL_INVALID_CELL_VALUE", "Non-finite number")
        return value
    raise _error("EXCEL_INVALID_CELL_VALUE", "Unsupported cell value")


def _headers(values: Sequence[Any]) -> tuple[str, ...]:
    items = list(values)
    while items and items[-1] is None:
        items.pop()
    if not items or any(not isinstance(x, str) or not x.strip() for x in items):
        raise _error("EXCEL_EMPTY_HEADER", "Every header must contain text")
    if len(items) > MAX_COLUMNS:
        raise _error("EXCEL_COLUMN_LIMIT_EXCEEDED", "Worksheet exceeds 500 columns")
    result = tuple(x.strip() for x in items)
    for value in result:
        _validate_text(value)
    if len(result) != len(set(result)):
        raise _error("EXCEL_DUPLICATE_HEADER", "Headers must be unique")
    return result


def _sheet_iterator(
    formulas: Any, cached: Any, index: int, deadline: float, budget: _Budget
) -> tuple[tuple[str, ...], Iterator[ExcelRow]]:
    source = formulas.worksheets[index]
    values = cached.worksheets[index]
    _reject_merges(formulas, source, deadline)
    source.reset_dimensions()
    values.reset_dimensions()
    source_rows = source.iter_rows()
    value_rows = values.iter_rows()
    try:
        source_header = next(source_rows)
        next(value_rows)
    except StopIteration as exc:
        source_rows.close()
        value_rows.close()
        raise _error("EXCEL_EMPTY_HEADER", "Worksheet must have a header") from exc
    except (
        ValueError,
        TypeError,
        KeyError,
        OSError,
        DefusedXmlException,
        ParseError,
    ) as exc:
        source_rows.close()
        value_rows.close()
        raise _error(
            "EXCEL_INVALID_WORKBOOK", "Worksheet cell data is invalid"
        ) from exc
    headers = _headers(tuple(x.value for x in source_header))
    budget.cells += len(headers)
    if budget.cells > MAX_CELLS:
        raise _error("EXCEL_CELL_LIMIT_EXCEEDED", "Workbook exceeds cell limit")

    def rows() -> Iterator[ExcelRow]:
        try:
            for row_number, (source_row, value_row) in enumerate(
                zip(source_rows, value_rows), start=2
            ):
                _check(deadline)
                if row_number - 1 > MAX_ROWS:
                    raise _error(
                        "EXCEL_ROW_LIMIT_EXCEEDED", "Worksheet exceeds row limit"
                    )
                width = max(len(source_row), len(value_row))
                budget.cells += width
                if width > MAX_COLUMNS:
                    raise _error(
                        "EXCEL_COLUMN_LIMIT_EXCEEDED", "Worksheet exceeds column limit"
                    )
                if budget.cells > MAX_CELLS:
                    raise _error(
                        "EXCEL_CELL_LIMIT_EXCEEDED", "Workbook exceeds cell limit"
                    )
                if any(x.value is not None for x in source_row[len(headers) :]):
                    raise _error("EXCEL_EMPTY_HEADER", "Data column has no header")
                output = []
                formula_columns = []
                for column in range(len(headers)):
                    source_cell = (
                        source_row[column] if column < len(source_row) else None
                    )
                    value_cell = value_row[column] if column < len(value_row) else None
                    value = getattr(value_cell, "value", None)
                    if getattr(source_cell, "data_type", None) == "f":
                        formula_columns.append(column)
                        if value is None:
                            raise _error(
                                "EXCEL_FORMULA_CACHE_MISSING",
                                "Formula has no saved value",
                            )
                    output.append(
                        _scalar(value, getattr(value_cell, "number_format", "General"))
                    )
                yield ExcelRow(tuple(output), tuple(formula_columns), row_number)
        except ProjectError:
            raise
        except (
            ValueError,
            TypeError,
            KeyError,
            OSError,
            DefusedXmlException,
            ParseError,
        ) as exc:
            raise _error(
                "EXCEL_INVALID_WORKBOOK", "Worksheet cell data is invalid"
            ) from exc
        finally:
            source_rows.close()
            value_rows.close()

    return headers, rows()


def _open(content: bytes) -> tuple[Any, Any]:
    formulas = None
    formula_buffer = io.BytesIO(content)
    cached_buffer = io.BytesIO(content)
    try:
        formulas = load_workbook(
            formula_buffer, read_only=True, data_only=False, keep_links=False
        )
        cached = load_workbook(
            cached_buffer, read_only=True, data_only=True, keep_links=False
        )
        setattr(formulas, "_autoflow_buffer", formula_buffer)  # noqa: B010
        setattr(cached, "_autoflow_buffer", cached_buffer)  # noqa: B010
        return formulas, cached
    except (
        OSError,
        ValueError,
        KeyError,
        TypeError,
        zipfile.BadZipFile,
        DefusedXmlException,
        ParseError,
    ) as exc:
        if formulas is not None:
            formulas.close()
        formula_buffer.close()
        cached_buffer.close()
        raise _error("EXCEL_INVALID_WORKBOOK", "Excel workbook invalid") from exc


def _close_workbook(workbook: Any) -> None:
    buffer = getattr(workbook, "_autoflow_buffer", None)
    workbook.close()
    if buffer is not None:
        buffer.close()


def inspect_workbook(path: Path) -> WorkbookInspection:
    deadline = time.monotonic() + MAX_SECONDS
    content, fingerprint, source_stat = _read_file(path, deadline)
    _validate_archive(content, deadline)
    formulas = cached = None
    try:
        formulas, cached = _open(content)
        budget = _Budget()
        sheets = []
        for index, sheet in enumerate(formulas.worksheets):
            headers, rows = _sheet_iterator(formulas, cached, index, deadline, budget)
            sample: list[tuple[ScalarValue, ...]] = []
            counts = [0] * len(headers)
            row_count = ignored = 0
            for row in rows:
                if all(x is None for x in row.values):
                    ignored += 1
                    continue
                row_count += 1
                for column in row.formula_columns:
                    counts[column] += 1
                if len(sample) < SAMPLE_ROWS:
                    sample.append(row.values)
            sheets.append(
                SheetInspection(
                    str(index + 1),
                    sheet.title,
                    headers,
                    tuple(sample),
                    row_count,
                    ignored,
                    tuple(counts),
                    tuple(x == 0 for x in counts),
                )
            )
        _, final_fingerprint, final_stat = _read_file(path, deadline)
        if fingerprint != final_fingerprint or _identity(source_stat) != _identity(
            final_stat
        ):
            raise _error("EXCEL_SOURCE_CHANGED", "Excel file changed during inspection")
        return WorkbookInspection(fingerprint, path.name, tuple(sheets))
    finally:
        if formulas is not None:
            _close_workbook(formulas)
        if cached is not None:
            _close_workbook(cached)


class ExcelRowStream(Iterator[ExcelRow]):
    def __init__(self, path: Path, sheet_id: str, expected_fingerprint: str):
        self._path = path
        self._expected = expected_fingerprint
        self._deadline = time.monotonic() + MAX_SECONDS
        self.closed = False
        content, fingerprint, self._source_stat = _read_file(path, self._deadline)
        if fingerprint != expected_fingerprint:
            raise _error("EXCEL_SOURCE_CHANGED", "Excel file changed after inspection")
        _validate_archive(content, self._deadline)
        self._formulas, self._cached = _open(content)
        self._rows: Iterator[ExcelRow] = iter(())
        try:
            index = int(sheet_id) - 1
            if (
                index < 0
                or index >= len(self._formulas.worksheets)
                or str(index + 1) != sheet_id
            ):
                raise ValueError
            _, self._rows = _sheet_iterator(
                self._formulas, self._cached, index, self._deadline, _Budget()
            )
        except (ValueError, IndexError) as exc:
            self.close()
            raise _error("EXCEL_SHEET_NOT_FOUND", "Worksheet not found") from exc
        except Exception:
            self.close()
            raise

    def __iter__(self) -> Self:
        return self

    def __next__(self) -> ExcelRow:
        if self.closed:
            raise StopIteration
        try:
            return next(self._rows)
        except StopIteration:
            try:
                _, fingerprint, final_stat = _read_file(self._path, self._deadline)
                if fingerprint != self._expected or _identity(final_stat) != _identity(
                    self._source_stat
                ):
                    raise _error(
                        "EXCEL_SOURCE_CHANGED", "Excel changed while rows read"
                    )
            finally:
                self.close()
            raise
        except ProjectError:
            self.close()
            raise
        except (
            ValueError,
            TypeError,
            KeyError,
            OSError,
            DefusedXmlException,
            ParseError,
        ) as exc:
            self.close()
            raise _error(
                "EXCEL_INVALID_WORKBOOK", "Worksheet cell data is invalid"
            ) from exc

    def __enter__(self) -> Self:
        return self

    def __exit__(self, *_args: object) -> None:
        self.close()

    def close(self) -> None:
        if not self.closed:
            self.closed = True
            close_rows = getattr(self._rows, "close", None)
            if close_rows is not None:
                close_rows()
            _close_workbook(self._formulas)
            _close_workbook(self._cached)


def read_sheet(path: Path, sheet_id: str, expected_fingerprint: str) -> ExcelRowStream:
    return ExcelRowStream(path, sheet_id, expected_fingerprint)


def _export(value: ScalarValue) -> tuple[Any, str | None]:
    if not isinstance(value, dict):
        if isinstance(value, float) and not math.isfinite(value):
            raise _error("EXCEL_INVALID_CELL_VALUE", "Non-finite number")
        if isinstance(value, str):
            _validate_text(value)
        return value, None
    raw = value.get("value")
    precision = value.get("precision")
    offset = value.get("offset")
    if (
        value.get("kind") != "date"
        or not isinstance(raw, str)
        or (offset is not None and not isinstance(offset, str))
    ):
        raise _error("EXCEL_INVALID_CELL_VALUE", "Invalid date scalar")
    try:
        if precision == "date":
            return datetime.combine(
                date.fromisoformat(raw), datetime.min.time()
            ), "yyyy-mm-dd"
        if precision == "datetime":
            parsed = datetime.fromisoformat(raw)
            if parsed.tzinfo is not None:
                raise ValueError
            if offset is not None:
                datetime.fromisoformat(
                    f"2000-01-01T00:00:00{offset.replace('Z', '+00:00')}"
                )
            result = raw + (offset or "")
            _validate_text(result)
            return result, "text"
    except ValueError as exc:
        raise _error("EXCEL_INVALID_CELL_VALUE", "Invalid date scalar") from exc
    raise _error("EXCEL_INVALID_CELL_VALUE", "Invalid date scalar")


def _validate_text(value: str) -> None:
    if len(value.encode("utf-16-le")) // 2 > 32_767:
        raise _error("EXCEL_STRING_TOO_LONG", "Excel text exceeds 32,767 UTF-16 units")


def _cleanup_writer_files(workbook: Any, existing_files: set[str]) -> None:
    for sheet in workbook.worksheets:
        writer = getattr(sheet, "_writer", None)
        if writer is None:
            continue
        output = getattr(writer, "out", None)
        if output in existing_files or output not in ALL_TEMP_FILES:
            continue
        try:
            sheet.close()
        except (ValueError, GeneratorExit):
            pass
        if output in ALL_TEMP_FILES:
            try:
                writer.cleanup()
            except OSError:
                if output in ALL_TEMP_FILES:
                    ALL_TEMP_FILES.remove(output)


def write_workbook(
    path: Path, headers: Sequence[str], rows: Iterable[Sequence[ScalarValue]]
) -> None:
    deadline = time.monotonic() + MAX_SECONDS
    if path.suffix.lower() != ".xlsx":
        raise _error("EXCEL_INVALID_OUTPUT", "Output must end in .xlsx")
    if path.exists() or path.is_symlink():
        raise _error("EXCEL_OUTPUT_EXISTS", "Output exists")
    names = _headers(tuple(headers))
    temporary = None
    workbook = Workbook(write_only=True)
    writer_files_before = set(ALL_TEMP_FILES)
    try:
        sheet = workbook.create_sheet("Data")
        header_cells = []
        for name in names:
            cell = WriteOnlyCell(sheet, value=name)
            cell.data_type = "s"
            header_cells.append(cell)
        sheet.append(header_cells)
        cells = len(names)
        for row_count, row in enumerate(rows, start=1):
            _check(deadline)
            if row_count > MAX_ROWS:
                raise _error("EXCEL_ROW_LIMIT_EXCEEDED", "Export exceeds row limit")
            values = tuple(row)
            if len(values) != len(names):
                raise _error("EXCEL_INVALID_ROW", "Row does not match headers")
            cells += len(values)
            if cells > MAX_CELLS:
                raise _error("EXCEL_CELL_LIMIT_EXCEEDED", "Export exceeds cell limit")
            output = []
            for value in values:
                converted, formatting = _export(value)
                cell = WriteOnlyCell(sheet, value=converted)
                if isinstance(value, str) or formatting == "text":
                    cell.data_type = "s"
                elif formatting:
                    cell.number_format = formatting
                output.append(cell)
            sheet.append(output)
        descriptor, raw = tempfile.mkstemp(
            prefix=f".{path.name}.", suffix=".tmp", dir=path.parent
        )
        os.close(descriptor)
        temporary = Path(raw)
        workbook.save(temporary)
        _check(deadline)
        if temporary.stat().st_size > MAX_FILE_BYTES:
            raise _error("EXCEL_FILE_TOO_LARGE", "Export exceeds 64 MiB")
        try:
            os.link(temporary, path)
        except FileExistsError as exc:
            raise _error("EXCEL_OUTPUT_EXISTS", "Output exists") from exc
    except ProjectError:
        raise
    except OSError as exc:
        raise _error("EXCEL_OUTPUT_UNAVAILABLE", "Export could not be written") from exc
    finally:
        _cleanup_writer_files(workbook, writer_files_before)
        workbook.close()
        if temporary is not None:
            try:
                temporary.unlink(missing_ok=True)
            except OSError:
                pass
