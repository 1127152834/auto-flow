"""Bounded XLSX rendering for workflow table exports.

Source behavior: WebRPA DataCollector._write_styled_excel at
reference/WebRPA@5ccb900e8dcf1530aae66f676d87593c416c7ebb.
"""

from __future__ import annotations

import asyncio
import io
import json
from collections.abc import Mapping, Sequence
from datetime import date, datetime, time
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from autoflow.domain.workflows.execution import CancellationToken
from autoflow.infrastructure.filesystem.project_excel import validate_workbook_content

_MAX_COLUMNS = 500
_MAX_CELLS = 1_000_000
_MAX_GRID_ROWS = 100_000
_MAX_NORMALIZED_BYTES = 32 * 1024 * 1024
_MAX_OUTPUT_BYTES = 64 * 1024 * 1024
_LARGE_DATA_THRESHOLD = 500


def _raise_if_cancelled(cancellation: CancellationToken | None) -> None:
    if cancellation is not None:
        cancellation.raise_if_cancelled()


def _normalize(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date, time)):
        return value
    if isinstance(value, (list, dict, tuple, set)):
        try:
            return json.dumps(value, ensure_ascii=False, default=str)
        except Exception:  # noqa: BLE001 -- preserve frozen fallback.
            return str(value)
    return str(value)


def _force_literal_string(cell: Any, value: Any) -> None:
    if isinstance(value, str):
        cell.data_type = "s"


def _render(
    rows: Sequence[Mapping[str, Any]],
    sheet_name: str,
    existing_content: bytes | None,
    cancellation: CancellationToken | None,
) -> bytes:
    _raise_if_cancelled(cancellation)
    columns: list[str] = []
    for row_index, row in enumerate(rows):
        if row_index % 256 == 0:
            _raise_if_cancelled(cancellation)
        for key in row:
            if key not in columns:
                columns.append(key)
                if len(columns) > _MAX_COLUMNS:
                    raise ValueError("Excel导出列数超过工作流安全限制")

    if len(rows) * len(columns) > _MAX_CELLS:
        raise ValueError("Excel导出单元格数量超过工作流安全限制")

    sampled_widths = [len(str(column)) for column in columns]
    normalized_bytes = 0
    for row_index, row in enumerate(rows):
        if row_index % 256 == 0:
            _raise_if_cancelled(cancellation)
        for column_index, column in enumerate(columns):
            value = _normalize(row.get(column))
            if value is not None:
                normalized_bytes += len(str(value).encode("utf-8"))
                if normalized_bytes > _MAX_NORMALIZED_BYTES:
                    raise ValueError("Excel导出内容超过工作流安全限制")
            if row_index < 200:
                sampled_widths[column_index] = max(
                    sampled_widths[column_index],
                    len(str(value)) if value is not None else 0,
                )

    existing_workbook = False
    if existing_content:
        try:
            validate_workbook_content(existing_content)
            probe = load_workbook(
                io.BytesIO(existing_content), read_only=True, data_only=False
            )
            try:
                existing_cells = 0
                for existing_sheet in probe.worksheets:
                    max_column = existing_sheet.max_column or 0
                    max_row = existing_sheet.max_row or 0
                    if max_column > _MAX_COLUMNS:
                        raise ValueError(
                            "已有Excel文件列数超过工作流安全限制"
                        )
                    existing_cells += max_row * max_column
                    if existing_cells > _MAX_CELLS:
                        raise ValueError(
                            "已有Excel文件单元格数量超过工作流安全限制"
                        )
            finally:
                probe.close()
            workbook = load_workbook(io.BytesIO(existing_content))
            existing_workbook = True
        except ValueError:
            raise
        except Exception as error:
            raise ValueError("已有Excel文件无效或超过安全限制") from error
    else:
        workbook = Workbook(write_only=True)
    if existing_workbook:
        if sheet_name in workbook.sheetnames:
            del workbook[sheet_name]
        sheet = workbook.create_sheet(sheet_name)
    else:
        sheet = workbook.create_sheet(sheet_name)
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(
        start_color="4472C4", end_color="4472C4", fill_type="solid"
    )
    header_alignment = Alignment(horizontal="center", vertical="center")
    header_border = Border(
        left=Side(style="thin", color="2F5496"),
        right=Side(style="thin", color="2F5496"),
        top=Side(style="thin", color="2F5496"),
        bottom=Side(style="thin", color="2F5496"),
    )
    for column_index, max_length in enumerate(sampled_widths, 1):
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            max_length + 4, 50
        )
    sheet.row_dimensions[1].height = 25
    sheet.freeze_panes = "A2"

    header_cells: list[Any] = []
    for column in columns:
        cell = WriteOnlyCell(sheet, value=column) if not existing_workbook else None
        if existing_workbook:
            cell = sheet.cell(row=1, column=len(header_cells) + 1, value=column)
        assert cell is not None
        _force_literal_string(cell, column)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = header_border
        header_cells.append(cell)
    if not existing_workbook:
        sheet.append(header_cells)

    is_large = len(rows) >= _LARGE_DATA_THRESHOLD
    even_fill = PatternFill(
        start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"
    )
    cell_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    output = io.BytesIO()
    try:
        for row_index, row in enumerate(rows):
            if row_index % 256 == 0:
                _raise_if_cancelled(cancellation)
            normalized = [_normalize(row.get(column)) for column in columns]
            if existing_workbook:
                sheet.append(normalized)
                excel_row = row_index + 2
                for data_cell in sheet[excel_row]:
                    _force_literal_string(data_cell, data_cell.value)
                    if not is_large:
                        if (row_index + 1) % 2 == 1:
                            data_cell.fill = even_fill
                        data_cell.border = cell_border
            else:
                cells = []
                for value in normalized:
                    data_cell = WriteOnlyCell(sheet, value=value)
                    _force_literal_string(data_cell, value)
                    if not is_large:
                        if row_index % 2 == 1:
                            data_cell.fill = even_fill
                        data_cell.border = cell_border
                    cells.append(data_cell)
                sheet.append(cells)
        workbook.save(output)
    finally:
        workbook.close()
    content = output.getvalue()
    if not content:
        raise OSError("Excel导出文件大小为 0")
    if len(content) > _MAX_OUTPUT_BYTES:
        raise ValueError("Excel导出内容超过工作流安全限制")
    _raise_if_cancelled(cancellation)
    return content


def _render_grid(
    rows: Sequence[Sequence[str]],
    sheet_name: str,
    header_row: int,
    include_header: bool,
    cancellation: CancellationToken | None,
) -> bytes:
    _raise_if_cancelled(cancellation)
    if len(rows) > _MAX_GRID_ROWS:
        raise ValueError("表格行数超过工作流安全限制")

    max_columns = 0
    cell_count = 0
    normalized_bytes = 0
    widths: list[int] = []
    for row_index, row in enumerate(rows):
        if row_index % 256 == 0:
            _raise_if_cancelled(cancellation)
        if isinstance(row, (str, bytes, bytearray)):
            raise TypeError("表格数据格式无效")
        max_columns = max(max_columns, len(row))
        if max_columns > _MAX_COLUMNS:
            raise ValueError("表格列数超过工作流安全限制")
        cell_count += len(row)
        if cell_count > _MAX_CELLS:
            raise ValueError("表格单元格数量超过工作流安全限制")
        while len(widths) < len(row):
            widths.append(0)
        for column_index, value in enumerate(row):
            normalized_bytes += len(value.encode("utf-8"))
            if normalized_bytes > _MAX_NORMALIZED_BYTES:
                raise ValueError("表格文本大小超过工作流安全限制")
            widths[column_index] = max(widths[column_index], len(value))

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(sheet_name)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(
        start_color="4472C4", end_color="4472C4", fill_type="solid"
    )
    header_alignment = Alignment(horizontal="center", vertical="center")
    body_alignment = Alignment(horizontal="left", vertical="center")
    for column_index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            width + 2, 50
        )

    output = io.BytesIO()
    try:
        for row_index, row in enumerate(rows):
            if row_index % 256 == 0:
                _raise_if_cancelled(cancellation)
            rendered_row: list[Any] = []
            for value in row:
                cell = WriteOnlyCell(sheet, value=value)
                _force_literal_string(cell, value)
                if include_header and row_index == header_row:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = header_alignment
                else:
                    cell.alignment = body_alignment
                rendered_row.append(cell)
            sheet.append(rendered_row)
        workbook.save(output)
    finally:
        workbook.close()
    content = output.getvalue()
    if not content:
        raise OSError("Excel导出文件大小为 0")
    if len(content) > _MAX_OUTPUT_BYTES:
        raise ValueError("Excel导出内容超过工作流安全限制")
    _raise_if_cancelled(cancellation)
    return content


class OpenpyxlTableWorkbookRenderer:
    """Render the frozen styled table contract without filesystem access."""

    async def render(
        self,
        *,
        rows: Sequence[Mapping[str, Any]],
        sheet_name: str,
        existing_content: bytes | None,
        cancellation: CancellationToken | None,
    ) -> bytes:
        return await asyncio.to_thread(
            _render, rows, sheet_name, existing_content, cancellation
        )

    async def render_grid(
        self,
        *,
        rows: Sequence[Sequence[str]],
        sheet_name: str,
        header_row: int,
        include_header: bool,
        cancellation: CancellationToken | None,
    ) -> bytes:
        return await asyncio.to_thread(
            _render_grid,
            rows,
            sheet_name,
            header_row,
            include_header,
            cancellation,
        )
