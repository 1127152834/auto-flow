from __future__ import annotations

import base64
import binascii
import hashlib
import io
import os
import re
import tempfile
import threading
import zipfile
from copy import deepcopy
from pathlib import Path, PurePosixPath
from typing import Any

import openpyxl
import xlrd  # type: ignore[import-untyped]
from docx import Document
from pypdf import PdfReader

from autoflow.domain.workflows.runs import WorkflowRunError
from autoflow.infrastructure.filesystem.project_excel import validate_workbook_content

MAX_ATTACHMENT_BYTES = 16 * 1024 * 1024
MAX_IMAGE_TOTAL_BYTES = 64 * 1024 * 1024
MAX_AUDIO_BYTES = 32 * 1024 * 1024
MAX_EXTRACTED_CHARS = 30_000
MAX_ARCHIVE_ENTRIES = 10_000
MAX_ARCHIVE_BYTES = 128 * 1024 * 1024
_IMAGE_TYPES = {
    "image/png": "png",
    "image/jpeg": "jpg",
    "image/webp": "webp",
    "image/gif": "gif",
    "image/bmp": "bmp",
    "image/svg+xml": "svg",
}
_TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".html",
    ".htm",
    ".svg",
    ".json",
    ".log",
    ".xml",
    ".yaml",
    ".yml",
    ".ini",
    ".tsv",
}
_WHISPER_MODEL_SIZES = {
    "tiny",
    "base",
    "small",
    "medium",
    "large",
    "large-v2",
    "large-v3",
}


class AssistantFileStore:
    def __init__(self, root: Path) -> None:
        self._root = root.resolve()
        self._attachments = self._root / "attachments"
        self._whisper_models: dict[str, Any] = {}
        self._whisper_lock = threading.Lock()

    def store_images(self, images: list[str]) -> list[str]:
        stored: list[str] = []
        total = 0
        for image in images:
            if image.startswith("assistant-attachment://"):
                self._attachment_path(image)
                stored.append(image)
                continue
            mime_type, content = _decode_image(image)
            total += len(content)
            if len(content) > MAX_ATTACHMENT_BYTES or total > MAX_IMAGE_TOTAL_BYTES:
                raise WorkflowRunError(
                    "ASSISTANT_ATTACHMENT_TOO_LARGE", "小助手图片附件超过大小限制", 422
                )
            name = f"{hashlib.sha256(content).hexdigest()}.{_IMAGE_TYPES[mime_type]}"
            target = self._attachments / name
            self._write_once(target, content)
            stored.append(f"assistant-attachment://{name}")
        return stored

    def hydrate_model_messages(
        self, messages: list[dict[str, Any]]
    ) -> list[dict[str, Any]]:
        hydrated = deepcopy(messages)
        for message in hydrated:
            content = message.get("content")
            if not isinstance(content, list):
                continue
            for part in content:
                image = part.get("image_url") if isinstance(part, dict) else None
                if isinstance(image, dict):
                    url = image.get("url")
                    if isinstance(url, str):
                        image["url"] = self.resolve_image(url)
        return hydrated

    def public_messages(
        self, messages: tuple[dict[str, Any], ...]
    ) -> list[dict[str, Any]]:
        result = deepcopy(list(messages))
        for message in result:
            images = message.get("images")
            if isinstance(images, list):
                message["images"] = [
                    self.resolve_image(image) if isinstance(image, str) else image
                    for image in images
                ]
        return result

    def resolve_image(self, reference: str) -> str:
        if not reference.startswith("assistant-attachment://"):
            return reference
        path = self._attachment_path(reference)
        try:
            content = path.read_bytes()
        except OSError as error:
            raise WorkflowRunError(
                "ASSISTANT_ATTACHMENT_MISSING", "小助手图片附件不存在", 409
            ) from error
        mime_type = next(
            (
                mime
                for mime, extension in _IMAGE_TYPES.items()
                if path.suffix == f".{extension}"
            ),
            None,
        )
        if mime_type is None:
            raise WorkflowRunError(
                "ASSISTANT_ATTACHMENT_INVALID", "小助手图片附件格式无效", 409
            )
        return f"data:{mime_type};base64,{base64.b64encode(content).decode()}"

    def extract_file(self, filename: str, encoded: str) -> dict[str, Any]:
        try:
            content = _decode_base64(encoded)
        except (ValueError, binascii.Error) as error:
            return {"success": False, "text": "", "error": f"文件解码失败: {error}"}
        if len(content) > MAX_ATTACHMENT_BYTES:
            return {"success": False, "text": "", "error": "文件超过 16 MiB 限制"}
        extension = Path(filename.strip() or "file").suffix.lower()
        try:
            if extension in _TEXT_EXTENSIONS:
                text = _decode_text(content)
            elif extension == ".pdf":
                text = "\n".join(
                    page.extract_text() or ""
                    for page in PdfReader(io.BytesIO(content)).pages[:50]
                )
            elif extension == ".docx":
                _validate_zip(content)
                document = Document(io.BytesIO(content))
                parts = [paragraph.text for paragraph in document.paragraphs]
                parts.extend(
                    "\t".join(cell.text for cell in row.cells)
                    for table in document.tables
                    for row in table.rows
                )
                text = "\n".join(parts)
            elif extension == ".xlsx":
                validate_workbook_content(content)
                workbook = openpyxl.load_workbook(
                    io.BytesIO(content), read_only=True, data_only=True
                )
                try:
                    parts = []
                    for sheet in workbook.worksheets:
                        parts.append(f"# 工作表: {sheet.title}")
                        for index, row in enumerate(sheet.iter_rows(values_only=True)):
                            if index >= 500:
                                parts.append("…（行数过多，仅解析前 500 行）")
                                break
                            parts.append(
                                "\t".join(
                                    "" if value is None else str(value) for value in row
                                )
                            )
                    text = "\n".join(parts)
                finally:
                    workbook.close()
            elif extension == ".xls":
                workbook = xlrd.open_workbook(file_contents=content)
                parts = []
                for sheet in workbook.sheets():
                    parts.append(f"# 工作表: {sheet.name}")
                    for index in range(min(sheet.nrows, 500)):
                        parts.append(
                            "\t".join(str(value) for value in sheet.row_values(index))
                        )
                text = "\n".join(parts)
            elif extension == ".doc":
                return {
                    "success": False,
                    "text": "",
                    "error": "暂不支持旧版 .doc（请另存为 .docx 后再上传）",
                }
            else:
                text = _decode_text(content)
        except Exception as error:  # noqa: BLE001 - normalize third-party parser errors
            return {
                "success": False,
                "text": "",
                "error": f"提取失败: {type(error).__name__}: {str(error)[:200]}",
            }
        return {"success": True, "text": _clip(text), "error": ""}

    def transcribe_audio(
        self, encoded: str, language: str, model_size: str
    ) -> dict[str, Any]:
        if model_size not in _WHISPER_MODEL_SIZES:
            return {"success": False, "text": "", "error": "不支持的语音模型规格"}
        if language != "auto" and not re.fullmatch(
            r"[A-Za-z]{2,3}(?:-[A-Za-z]{2,8})?", language
        ):
            return {"success": False, "text": "", "error": "语音语言参数无效"}
        try:
            content = _decode_base64(encoded)
        except (ValueError, binascii.Error) as error:
            return {"success": False, "text": "", "error": f"音频解码失败: {error}"}
        if not content:
            return {"success": False, "text": "", "error": "空音频"}
        if len(content) > MAX_AUDIO_BYTES:
            return {"success": False, "text": "", "error": "音频超过 32 MiB 限制"}

        temporary_dir = self._root / "tmp"
        temporary_dir.mkdir(parents=True, exist_ok=True)
        path = ""
        try:
            with tempfile.NamedTemporaryFile(
                dir=temporary_dir, suffix=".webm", delete=False
            ) as output:
                output.write(content)
                path = output.name
            with self._whisper_lock:
                model = self._whisper_models.get(model_size)
                if model is None:
                    from faster_whisper import (  # type: ignore[import-untyped]
                        WhisperModel,
                    )

                    model = WhisperModel(
                        f"Systran/faster-whisper-{model_size}",
                        device="cpu",
                        compute_type="int8",
                        download_root=str(self._root / "whisper_models"),
                    )
                    self._whisper_models[model_size] = model
            segments, info = model.transcribe(
                path, language=None if language == "auto" else language
            )
            text = "".join(segment.text for segment in segments).strip()
            return {
                "success": True,
                "text": text,
                "language": getattr(info, "language", language),
                "error": "",
            }
        except Exception as error:  # noqa: BLE001 - normalize native decoder/model errors
            return {
                "success": False,
                "text": "",
                "error": f"识别失败: {type(error).__name__}: {str(error)[:200]}",
            }
        finally:
            if path:
                Path(path).unlink(missing_ok=True)

    def _attachment_path(self, reference: str) -> Path:
        name = reference.removeprefix("assistant-attachment://")
        if not name or PurePosixPath(name).name != name:
            raise WorkflowRunError(
                "ASSISTANT_ATTACHMENT_INVALID", "小助手图片附件标识无效", 422
            )
        path = (self._attachments / name).resolve()
        if not path.is_relative_to(self._attachments.resolve()):
            raise WorkflowRunError(
                "ASSISTANT_ATTACHMENT_INVALID", "小助手图片附件标识无效", 422
            )
        return path

    @staticmethod
    def _write_once(target: Path, content: bytes) -> None:
        target.parent.mkdir(parents=True, exist_ok=True)
        try:
            with target.open("xb") as output:
                output.write(content)
                output.flush()
                os.fsync(output.fileno())
        except FileExistsError:
            return


def _decode_base64(value: str) -> bytes:
    raw = (
        value.split(",", 1)[1]
        if value.strip().startswith("data:") and "," in value
        else value
    )
    return base64.b64decode(raw, validate=True)


def _decode_image(value: str) -> tuple[str, bytes]:
    if not value.startswith("data:") or ";base64," not in value:
        raise WorkflowRunError(
            "ASSISTANT_ATTACHMENT_INVALID", "小助手图片附件必须是 base64 data URL", 422
        )
    metadata, encoded = value[5:].split(",", 1)
    mime_type = metadata.removesuffix(";base64").lower()
    if mime_type not in _IMAGE_TYPES:
        raise WorkflowRunError(
            "ASSISTANT_ATTACHMENT_INVALID", "小助手图片附件格式不受支持", 422
        )
    try:
        return mime_type, base64.b64decode(encoded, validate=True)
    except (ValueError, binascii.Error):
        raise WorkflowRunError(
            "ASSISTANT_ATTACHMENT_INVALID", "小助手图片附件解码失败", 422
        ) from None


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gbk", "gb18030", "utf-16", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="ignore")


def _clip(text: str) -> str:
    if len(text) <= MAX_EXTRACTED_CHARS:
        return text
    return text[:MAX_EXTRACTED_CHARS] + "\n…（内容过长，已截断，仅保留前 30000 字）"


def _validate_zip(content: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_ARCHIVE_ENTRIES:
            raise ValueError("压缩文档条目过多")
        total = 0
        for entry in entries:
            path = PurePosixPath(entry.filename)
            if path.is_absolute() or ".." in path.parts or entry.flag_bits & 1:
                raise ValueError("压缩文档结构不安全")
            total += entry.file_size
            if total > MAX_ARCHIVE_BYTES:
                raise ValueError("压缩文档展开后过大")
